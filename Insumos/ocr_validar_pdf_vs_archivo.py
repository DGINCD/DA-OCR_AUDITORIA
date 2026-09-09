import argparse
from collections import Counter
from datetime import datetime
import json
import math
import re
import sqlite3
import sys
import unicodedata
from pathlib import Path

try:
    from rapidocr_onnxruntime import RapidOCR as _PreloadedRapidOCR
except ImportError:
    _PreloadedRapidOCR = None

import pandas as pd
from PIL import Image, ImageEnhance, ImageFilter


DEFAULT_ID_PATTERN = r"(?<!\d)\d{5,8}(?!\d)"
OCR_CACHE_VERSION = "v7"
FIELDS_TO_VALIDATE = [
    "NumeroCertificado",
    "ApellidoNombre",
    "DNI",
    "CUIL",
    "Sexo",
    "FechaNacimiento",
    "Diagnostico",
    "FuncionesCorporales",
    "EstructurasCorporales",
    "ActividadParticipacion",
    "FactoresAmbientales",
    "OrientacionPrestacional",
    "FechaEmision",
    "FechaVencimiento",
    "Acompaniante",
]
FIELD_LABELS = {
    "NumeroCertificado": "Numero Certificado",
    "ApellidoNombre": "Apellido Nombre",
    "DNI": "DNI",
    "CUIL": "CUIL",
    "Sexo": "Sexo",
    "FechaNacimiento": "Fecha Nacimiento",
    "Diagnostico": "Diagnostico",
    "FuncionesCorporales": "Funciones Corporales",
    "EstructurasCorporales": "Estructuras Corporales",
    "ActividadParticipacion": "Actividad Participacion",
    "FactoresAmbientales": "Factores Ambientales",
    "OrientacionPrestacional": "Orientacion Prestacional",
    "FechaEmision": "Fecha Emision",
    "FechaVencimiento": "Fecha Vencimiento",
    "Acompaniante": "Acompaniante",
}
OCR_FIELD_LABELS = {
    field: f"OCR - {label}" for field, label in FIELD_LABELS.items()
}
USER_REVIEW_COLUMNS = [
    "Auditado usuario",
    "Resultado usuario",
    "Observacion usuario",
]
CIF_FIELDS = {
    "FuncionesCorporales",
    "EstructurasCorporales",
    "ActividadParticipacion",
    "FactoresAmbientales",
}
CIF_PREFIX_BY_FIELD = {
    "FuncionesCorporales": "B",
    "EstructurasCorporales": "S",
    "ActividadParticipacion": "D",
    "FactoresAmbientales": "E",
}
LAYOUT_SECTION_FIELDS = [
    "NumeroCertificado",
    "Diagnostico",
    "FuncionesCorporales",
    "EstructurasCorporales",
    "ActividadParticipacion",
    "FactoresAmbientales",
    "OrientacionPrestacional",
    "Acompaniante",
]


def normalize_text(value):
    if pd.isna(value):
        return ""
    text = str(value).upper().strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"\s+", " ", text)
    return text


def compact_text(value):
    return re.sub(r"[^A-Z0-9]+", "", normalize_text(value))


def normalize_identifier(value):
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = normalize_text(value)
    text = re.sub(r"\.0$", "", text)
    return re.sub(r"\D", "", text)


def normalize_date_candidates(value):
    text = normalize_text(value)
    if not text:
        return []

    parsed = pd.to_datetime(value, errors="coerce", dayfirst=True)
    if not pd.isna(parsed):
        day = int(parsed.day)
        month = int(parsed.month)
        year = int(parsed.year)
        return [
            f"{day}/{month}/{year}",
            f"{day:02d}/{month:02d}/{year}",
            f"{year}{month:02d}{day:02d}",
            f"{month:02d}-{year}",
            f"{month}-{year}",
        ]

    return [text]


def certificate_components(value):
    text = normalize_text(value)
    compact = compact_text(text)
    numbers = re.findall(r"\d+", compact)
    dni = ""
    dates = []
    for number in numbers:
        if len(number) == 8 and not dni:
            dni = number
        elif len(number) == 8:
            dates.append(number)
    prefix = compact[:3] if len(compact) >= 3 else ""
    return {
        "compact": compact,
        "prefix": prefix,
        "dni": dni,
        "dates": dates[:2],
    }


def dni_from_certificate(value):
    text = normalize_text(value)
    if not text:
        return ""

    match = re.search(
        r"[A-Z]{3}\s*[- ]?\s*\d{2}\s*[- ]?\s*0*(\d{5,8})\s*[- ]?\s*(?:19|20)\d{6}",
        text,
    )
    if not match:
        return ""
    return normalize_identifier(match.group(1))


def recover_missing_identifiers_from_certificate(df_team, id_column):
    cert_column = certificate_column_name(df_team)
    if not cert_column:
        return 0, []

    recovered_count = 0
    mismatches = []
    for index, row in df_team.iterrows():
        identifier = normalize_identifier(row.get(id_column, ""))
        certificate_identifier = dni_from_certificate(row.get(cert_column, ""))
        if not certificate_identifier:
            continue
        if not identifier:
            df_team.at[index, id_column] = certificate_identifier
            df_team.at[index, "identificador"] = certificate_identifier
            recovered_count += 1
        elif identifier != certificate_identifier:
            mismatches.append(
                {
                    "fila_excel": index + 2,
                    "dni_columna": identifier,
                    "dni_certificado": certificate_identifier,
                    "certificado": row.get(cert_column, ""),
                }
            )
    return recovered_count, mismatches


def certificate_score(expected, ocr_text):
    components = certificate_components(expected)
    if not components["compact"]:
        return "NO_ENCONTRADO", 0, ""

    compact_ocr = compact_text(ocr_text)
    if components["compact"] in compact_ocr:
        return "OK", 1, "certificado completo"

    hits = 0
    checks = []
    if components["dni"]:
        checks.append(("dni", components["dni"] in compact_ocr))
    for idx, date_value in enumerate(components["dates"], start=1):
        checks.append((f"fecha{idx}", date_value in compact_ocr))
    # El prefijo puede deformarse por OCR; suma, pero no es requisito principal.
    if components["prefix"]:
        checks.append(("prefijo", components["prefix"] in compact_ocr))

    hits = sum(1 for _, found in checks if found)
    required = max(len(checks) - 1, 1)
    if hits >= required and components["dni"]:
        return "OK", round(hits / max(len(checks), 1), 3), "componentes certificado: " + ", ".join(name for name, found in checks if found)
    return "NO_ENCONTRADO", round(hits / max(len(checks), 1), 3), "componentes certificado insuficientes"


def code_tokens(value):
    normalized = normalize_text(value)
    normalized = re.sub(r"\.{2,}", ".", normalized)
    pattern = r"([A-Z]\d{2,5}\.\+?\d{1,3})(?=[A-Z]|[^A-Z0-9]|$)"
    tokens = []
    for match in re.finditer(pattern, normalized):
        token = match.group(1)
        previous_char = normalized[match.start() - 1] if match.start() > 0 else ""
        if previous_char.isalpha():
            continue
        previous_text = normalized[max(0, match.start() - 2) : match.start()]
        next_text = normalized[match.end() : match.end() + 1]
        if re.search(r"[A-Z]\.$", previous_text):
            continue
        if next_text and next_text.isalpha() and next_text not in CIF_PREFIX_BY_FIELD.values():
            continue
        tokens.append(token)
    return tokens


def code_set(value, prefix=None):
    codes = {code.upper() for code in code_tokens(value)}
    if prefix:
        codes = {code for code in codes if code.startswith(prefix)}
    return codes


def extract_cif_section_text(ocr_text, field):
    normalized = normalize_text(ocr_text)
    section_patterns = {
        "FuncionesCorporales": (
            r"FUNCIONES\s*CORPORALES[:\s]*",
            r"ESTRUCTURAS\s*CORPORALES|ACTIVIDAD\s*/?\s*PARTICIPACION|FACTORES\s*AMBIENTALES|ORIENTACION\s*PRESTACIONAL|ACTUALIZACION",
        ),
        "EstructurasCorporales": (
            r"ESTRUCTURAS\s*CORPORALES[:\s]*",
            r"ACTIVIDAD\s*/?\s*PARTICIPACION|FACTORES\s*AMBIENTALES|ORIENTACION\s*PRESTACIONAL|ACTUALIZACION",
        ),
        "ActividadParticipacion": (
            r"ACTIVIDAD\s*/?\s*PARTICIPACION[:\s]*",
            r"FACTORES\s*AMBIENTALES|ORIENTACION\s*PRESTACIONAL|ACTUALIZACION|ACOMPANANTE",
        ),
        "FactoresAmbientales": (
            r"FACTORES\s*AMBIENTALES[:\s]*",
            r"ORIENTACION\s*PRESTACIONAL|ACTUALIZACION|ACOMPANANTE|LUGAR\s*Y\s*FECHA",
        ),
    }
    start_pattern, end_pattern = section_patterns[field]
    chunks = []
    for match in re.finditer(start_pattern, normalized, flags=re.IGNORECASE):
        start = match.end()
        tail = normalized[start : start + 1800]
        end = re.search(end_pattern, tail, flags=re.IGNORECASE)
        chunks.append(tail[: end.start()] if end else tail)
    return "\n".join(chunks)


def cif_code_comparison(field, expected, ocr_text):
    prefix = CIF_PREFIX_BY_FIELD[field]
    expected_codes = code_set(expected, prefix=prefix)
    section_text = extract_cif_section_text(ocr_text, field)
    observed_codes = code_set(section_text, prefix=prefix)

    if not observed_codes:
        observed_codes = code_set(ocr_text, prefix=prefix)
        source = "texto_completo"
    else:
        source = "seccion"

    all_text_codes = code_set(ocr_text, prefix=prefix)
    observed_codes = observed_codes | (expected_codes & all_text_codes)
    compact_ocr = compact_text(ocr_text)
    observed_codes = observed_codes | {
        code for code in expected_codes if compact_text(code) in compact_ocr
    }

    missing_set, extra_set, variant_matches = remove_ocr_variant_overlaps(
        expected_codes,
        observed_codes,
    )
    exact_matches = expected_codes & observed_codes
    missing_in_pdf = sorted(missing_set)
    extra_in_pdf = sorted(extra_set)
    matching = sorted(exact_matches | variant_matches)
    total = max(len(expected_codes | observed_codes), 1)
    score = len(matching) / total
    return {
        "expected": expected_codes,
        "observed": observed_codes,
        "missing_in_pdf": missing_in_pdf,
        "extra_in_pdf": extra_in_pdf,
        "score": score,
        "source": source,
    }


def format_code_list(codes):
    return ", ".join(sorted(codes))


def cif_variant_map(expected_codes, observed_codes):
    variants = []
    exact_matches = set(expected_codes) & set(observed_codes)
    for expected_code in sorted(expected_codes):
        if expected_code in exact_matches:
            continue
        exact_expected = normalized_code_for_similarity(expected_code)
        for observed_code in sorted(observed_codes):
            if observed_code in exact_matches:
                continue
            if expected_code == observed_code:
                continue
            if codes_equivalent_or_ocr_variant(expected_code, observed_code):
                variants.append(f"{expected_code} ~= {observed_code}")
                continue
            observed_norm = normalized_code_for_similarity(observed_code)
            if exact_expected and observed_norm:
                if observed_norm.startswith(exact_expected[:4]) or exact_expected.startswith(observed_norm[:4]):
                    variants.append(f"{expected_code} ? {observed_code}")
    return sorted(set(variants))


def cif_detail_row(pdf_name, identifier, audit_key, field, expected, ocr_text):
    if field not in CIF_FIELDS:
        return None

    prefix = CIF_PREFIX_BY_FIELD[field]
    comparison = cif_code_comparison(field, expected, ocr_text)
    section_text = extract_cif_section_text(ocr_text, field)
    section_codes = code_set(section_text, prefix=prefix)
    text_codes = code_set(ocr_text, prefix=prefix)
    variants = cif_variant_map(comparison["expected"], text_codes)
    exact_matches = comparison["expected"] & comparison["observed"]
    if not comparison["missing_in_pdf"] and not comparison["extra_in_pdf"]:
        variants = []

    return {
        "pdf": pdf_name,
        "identificador": identifier,
        "clave_auditoria": audit_key,
        "campo": field,
        "campo_nombre": FIELD_LABELS.get(field, field),
        "fuente_validacion": comparison["source"],
        "codigos_esperados_archivo": format_code_list(comparison["expected"]),
        "codigos_leidos_seccion": format_code_list(section_codes),
        "codigos_leidos_texto_completo": format_code_list(text_codes),
        "codigos_usados_validacion": format_code_list(comparison["observed"]),
        "codigos_encontrados": format_code_list(exact_matches),
        "codigos_no_encontrados_pdf": format_code_list(comparison["missing_in_pdf"]),
        "codigos_certificado_no_presentes_archivo": format_code_list(comparison["extra_in_pdf"]),
        "posibles_variantes_ocr": " | ".join(variants),
        "score_cif": comparison["score"],
    }


def order_cif_detail_columns(df):
    output = df.copy()
    if output.empty:
        return output

    def clean_value(value):
        if pd.isna(value):
            return ""
        text = str(value).strip()
        return "" if text.lower() == "nan" else text

    def best_ocr_value(row):
        validation_value = clean_value(row.get("codigos_usados_validacion", ""))
        if validation_value:
            return validation_value
        section_value = clean_value(row.get("codigos_leidos_seccion", ""))
        if section_value:
            return section_value
        text_value = clean_value(row.get("codigos_leidos_texto_completo", ""))
        if text_value:
            return text_value
        return clean_value(row.get("Datos OCR", ""))

    def missing_summary(row):
        status = clean_value(row.get("estado", ""))
        if status == "OK":
            return ""
        parts = []
        missing_pdf = clean_value(row.get("codigos_no_encontrados_pdf", ""))
        missing_file = clean_value(row.get("codigos_certificado_no_presentes_archivo", ""))
        variants = clean_value(row.get("posibles_variantes_ocr", ""))
        if missing_pdf:
            parts.append(f"No encontrado en PDF: {missing_pdf}")
        if missing_file:
            parts.append(f"Falta en archivo: {missing_file}")
        if variants:
            parts.append(f"Posibles variantes OCR: {variants}")
        if parts:
            return " | ".join(parts)
        return clean_value(row.get("Faltantes", ""))

    output["Datos archivo"] = output.apply(
        lambda row: clean_value(row.get("codigos_esperados_archivo", ""))
        or clean_value(row.get("Datos archivo", "")),
        axis=1,
    )
    output["Datos OCR"] = output.apply(best_ocr_value, axis=1)
    output["Faltantes"] = output.apply(missing_summary, axis=1)

    preferred = [
        "ID Auditoria",
        "DNI Auditoria",
        "pdf",
        "campo",
        "campo_nombre",
        "Datos archivo",
        "Datos OCR",
        "Faltantes",
    ]
    return output[[column for column in preferred if column in output.columns]]


def normalized_code_for_similarity(code):
    return re.sub(r"[^A-Z0-9]", "", normalize_text(code))


def codes_equivalent_or_ocr_variant(left_code, right_code):
    left = normalized_code_for_similarity(left_code)
    right = normalized_code_for_similarity(right_code)
    if not left or not right:
        return False
    if left == right:
        return True
    shorter, longer = sorted([left, right], key=len)
    # OCR often drops the last qualifier digit: D240.23 -> D240.2.
    if len(longer) - len(shorter) == 1 and longer.startswith(shorter):
        return True
    return False


def remove_ocr_variant_overlaps(expected_codes, observed_codes):
    expected = set(expected_codes)
    observed = set(observed_codes)
    matched_expected = set()
    matched_observed = set()
    for expected_code in expected:
        for observed_code in observed:
            if codes_equivalent_or_ocr_variant(expected_code, observed_code):
                matched_expected.add(expected_code)
                matched_observed.add(observed_code)
    return expected - matched_expected, observed - matched_observed, matched_expected


def similar_codes(missing_code, ocr_text):
    target = normalized_code_for_similarity(missing_code)
    if not target:
        return []

    candidates = sorted(set(code_tokens(ocr_text)))
    similar = []
    for candidate in candidates:
        candidate_norm = normalized_code_for_similarity(candidate)
        if not candidate_norm:
            continue
        if candidate_norm.startswith(target[:4]) or target.startswith(candidate_norm[:4]):
            similar.append(candidate)
    return similar[:5]


def word_tokens(value):
    tokens = re.findall(r"[A-Z0-9]+", normalize_text(value))
    ignored = {"DE", "DEL", "LA", "EL", "Y", "A", "EN", "CON", "LOS", "LAS"}
    return [token for token in tokens if len(token) >= 3 and token not in ignored]


def diagnosis_key_tokens(value):
    tokens = word_tokens(value)
    ignored = {
        "OTROS",
        "OTRAS",
        "PARTE",
        "PARTES",
        "SIN",
        "LOS",
        "LAS",
        "DEL",
        "CON",
        "POR",
        "NO",
        "ESPECIFICADO",
        "ESPECIFICADA",
        "CLASIFICADOS",
        "CLASIFICADAS",
    }
    return [token for token in tokens if token not in ignored and len(token) >= 4]


def field_score(field, expected, ocr_text):
    expected_text = normalize_text(expected)
    if field == "CUIL" and not expected_text:
        return "OK", 1, "CUIL/CUIT no informado en origen; campo variable por formato de certificado"
    if not expected_text:
        return "SIN_DATO_ORIGEN", 0, ""

    if field in CIF_FIELDS:
        comparison = cif_code_comparison(field, expected, ocr_text)
        if not comparison["observed"]:
            return "NO_ENCONTRADO", 0, "sin codigos OCR en seccion"
        coverage = len(comparison["expected"] & comparison["observed"]) / max(len(comparison["expected"]), 1)
        if coverage < 0.6:
            return "NO_ENCONTRADO", round(coverage, 3), (
                f"baja cobertura OCR: {len(comparison['expected'] & comparison['observed'])}/"
                f"{len(comparison['expected'])} codigos esperados"
            )
        if comparison["missing_in_pdf"] or comparison["extra_in_pdf"]:
            return "PARCIAL", round(comparison["score"], 3), (
                f"{len(comparison['expected'] & comparison['observed'])}/"
                f"{len(comparison['expected'] | comparison['observed'])} codigos ({comparison['source']})"
            )
        return "OK", 1, f"{len(comparison['expected'])} codigos ({comparison['source']})"

    compact_ocr = compact_text(ocr_text)
    compact_expected = compact_text(expected_text)

    if field == "NumeroCertificado":
        return certificate_score(expected, ocr_text)

    if field == "DNI":
        digits = normalize_identifier(expected)
        found = bool(digits and digits in re.sub(r"\D", "", ocr_text))
        return ("OK" if found else "NO_ENCONTRADO", 1 if found else 0, digits)

    if field == "CUIL":
        digits = normalize_identifier(expected)
        found = bool(digits and digits in re.sub(r"\D", "", ocr_text))
        if found:
            return "OK", 1, digits
        return (
            "OK",
            0.5,
            "CUIL/CUIT no localizado; puede no estar presente por formato de certificado",
        )

    if field in {"FechaNacimiento", "FechaEmision", "FechaVencimiento"}:
        candidates = normalize_date_candidates(expected)
        for candidate in candidates:
            if compact_text(candidate) in compact_ocr:
                return "OK", 1, " | ".join(candidates)
        return "NO_ENCONTRADO", 0, " | ".join(candidates)

    if field == "Sexo":
        expected_norm = normalize_text(expected)
        compact = compact_text(ocr_text)
        if expected_norm.startswith("M"):
            found = bool(re.search(r"DN[I1]M\d", compact))
            return ("OK" if found else "NO_ENCONTRADO", 1 if found else 0, "M")
        if expected_norm.startswith("F"):
            found = bool(re.search(r"DN[I1]F\d", compact))
            return ("OK" if found else "NO_ENCONTRADO", 1 if found else 0, "F")
        return "NO_ENCONTRADO", 0, expected_norm

    if field == "Acompaniante":
        expected_norm = normalize_text(expected)
        expected_short = "SI" if expected_norm.startswith("S") else "NO"
        normalized_ocr = normalize_text(ocr_text)
        around = re.search(r"ACOMPANANTE(.{0,250})", normalized_ocr)
        found = expected_short in compact_text(around.group(1)) if around else expected_short in compact_ocr
        return ("OK" if found else "NO_ENCONTRADO", 1 if found else 0, expected_short)

    if compact_expected and compact_expected in compact_ocr:
        return "OK", 1, compact_expected

    tokens = code_tokens(expected_text) if field not in {"ApellidoNombre", "Diagnostico", "OrientacionPrestacional"} else []
    if not tokens:
        tokens = word_tokens(expected_text)

    if not tokens:
        return "NO_ENCONTRADO", 0, compact_expected

    hits = sum(1 for token in tokens if compact_text(token) in compact_ocr)
    score = hits / len(tokens)
    if field == "ApellidoNombre":
        # En certificados escaneados el OCR suele unir apellidos y deformar nombres de pila.
        # Si al menos los dos primeros tokens del nombre aparecen, alcanza para validar identidad.
        key_tokens = tokens[:2] if len(tokens) >= 2 else tokens
        key_hits = sum(1 for token in key_tokens if compact_text(token) in compact_ocr)
        if key_tokens and key_hits == len(key_tokens):
            return "OK", round(max(score, key_hits / len(key_tokens)), 3), f"{key_hits}/{len(key_tokens)} tokens clave"

    if field == "Diagnostico":
        key_tokens = diagnosis_key_tokens(expected_text)
        key_hits = sum(1 for token in key_tokens if compact_text(token) in compact_ocr)
        key_score = key_hits / max(len(key_tokens), 1)
        if key_tokens and key_hits >= 3 and key_score >= 0.5:
            return "OK", round(max(score, key_score), 3), f"{key_hits}/{len(key_tokens)} tokens diagnostico"

    ok_threshold = 0.6 if field == "Diagnostico" else 0.85
    partial_threshold = 0.45 if field == "Diagnostico" else 0.55
    if score >= ok_threshold:
        status = "OK"
    elif score >= partial_threshold:
        status = "PARCIAL"
    else:
        status = "NO_ENCONTRADO"
    return status, round(score, 3), f"{hits}/{len(tokens)} tokens"


def missing_detail(field, expected, ocr_text):
    if field in CIF_FIELDS:
        comparison = cif_code_comparison(field, expected, ocr_text)
        coverage = len(comparison["expected"] & comparison["observed"]) / max(len(comparison["expected"]), 1)
        if not comparison["observed"] or coverage < 0.6:
            return (
                "No validable por OCR: no se leyo la seccion de "
                + FIELD_LABELS.get(field, field)
                + " con claridad"
            )
        parts = []
        if comparison["missing_in_pdf"]:
            missing_parts = []
            for token in comparison["missing_in_pdf"]:
                candidates = similar_codes(token, ocr_text)
                if candidates:
                    missing_parts.append(f"{token} (posible en PDF: {', '.join(candidates)})")
                else:
                    missing_parts.append(token)
            parts.append(
                "Codigo del archivo no encontrado en certificado: "
                + ", ".join(missing_parts)
            )
        if comparison["extra_in_pdf"]:
            parts.append(
                "Codigo del certificado no presente en archivo: "
                + ", ".join(comparison["extra_in_pdf"])
            )
        return " | ".join(parts) if parts else "Codigos CIF no coinciden con claridad"

    if field == "NumeroCertificado":
        return "No encontrado en PDF: no se leyo el numero de certificado completo"

    if field == "Diagnostico":
        return "No encontrado en PDF: no se leyo el diagnostico con claridad"

    if field == "DNI":
        return "No encontrado en PDF: no se leyo el numero esperado"

    if field == "CUIL":
        return "No validable por OCR: CUIL/CUIT puede no estar impreso en este formato de certificado"

    if field.startswith("Fecha"):
        return "No encontrado en PDF: no se leyo la fecha esperada"

    return "No encontrado en PDF: no se encontro coincidencia clara"


def extra_cif_detail(field, expected, ocr_text):
    if field not in CIF_FIELDS:
        return ""
    comparison = cif_code_comparison(field, expected, ocr_text)
    if not comparison["extra_in_pdf"]:
        return ""
    return (
        f"{FIELD_LABELS.get(field, field)}: Codigo del certificado no presente en archivo: "
        + ", ".join(comparison["extra_in_pdf"])
    )


def is_ocr_uncertain_alert(row):
    if row.estado == "SIN_DATO_ORIGEN":
        return False
    detail = str(row.detalle_revision)
    return detail.startswith("No validable por OCR")


def read_team_file(path, sheet_name=None):
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xls"}:
        if sheet_name:
            return pd.read_excel(path, sheet_name=sheet_name)
        workbook = pd.ExcelFile(path)
        selected_sheet = "SECLYT" if "SECLYT" in workbook.sheet_names else workbook.sheet_names[0]
        return pd.read_excel(path, sheet_name=selected_sheet)
    if suffix == ".csv":
        return pd.read_csv(path)
    raise ValueError(f"Formato no soportado: {suffix}. Use .xlsx, .xls o .csv")


def render_pdf_pages(pdf_path, dpi, max_pages=None):
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError("Falta PyMuPDF. Instale con: pip install pymupdf") from exc

    doc = fitz.open(str(pdf_path))
    total_pages = len(doc) if max_pages is None else min(len(doc), max_pages)
    zoom = dpi / 72
    matrix = fitz.Matrix(zoom, zoom)

    for page_index in range(total_pages):
        page = doc[page_index]
        pixmap = page.get_pixmap(matrix=matrix, alpha=False)
        image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
        yield page_index + 1, image


def ocr_pdf_tesseract(pdf_path, dpi, lang, tesseract_cmd=None, max_pages=None):
    try:
        import pytesseract
    except ImportError as exc:
        raise RuntimeError("Falta pytesseract. Instale con: pip install pytesseract") from exc

    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd

    page_texts = []
    for page_number, image in render_pdf_pages(pdf_path, dpi=dpi, max_pages=max_pages):
        text = pytesseract.image_to_string(image, lang=lang)
        page_texts.append({"pagina": page_number, "texto": text})
    return page_texts


def ocr_pdf_rapidocr(pdf_path, dpi, max_pages=None):
    try:
        import numpy as np
        RapidOCR = _PreloadedRapidOCR
        if RapidOCR is None:
            from rapidocr_onnxruntime import RapidOCR
    except ImportError as exc:
        raise RuntimeError(
            f"Falta RapidOCR o una dependencia binaria no cargo correctamente: {exc}"
        ) from exc

    engine = RapidOCR()
    page_texts = []
    for page_number, image in render_pdf_pages(pdf_path, dpi=dpi, max_pages=max_pages):
        result, _ = engine(np.array(image))
        lines = []
        for item in result or []:
            if len(item) >= 2:
                lines.append(str(item[1]))
        page_texts.append({"pagina": page_number, "texto": "\n".join(lines)})
    return page_texts


def otsu_image(image):
    import cv2
    import numpy as np

    gray = np.array(image.convert("L"))
    denoised = cv2.fastNlMeansDenoising(gray, None, 18, 7, 21)
    thresholded = cv2.threshold(denoised, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]
    return Image.fromarray(thresholded)


def enhanced_ocr_image(image):
    enhanced = image.convert("L")
    enhanced = ImageEnhance.Contrast(enhanced).enhance(1.8)
    enhanced = ImageEnhance.Sharpness(enhanced).enhance(2.0)
    enhanced = enhanced.filter(ImageFilter.SHARPEN)
    return enhanced.convert("RGB")


def rapidocr_text(engine, image):
    import numpy as np

    result, _ = engine(np.array(image.convert("RGB")))
    lines = []
    for item in result or []:
        if len(item) >= 2:
            lines.append(str(item[1]))
    return "\n".join(lines)


def rapidocr_items_with_boxes(engine, image):
    import numpy as np

    result, _ = engine(np.array(image.convert("RGB")))
    items = []
    for item in result or []:
        if len(item) < 2:
            continue
        box = item[0]
        text = str(item[1])
        if not text.strip():
            continue
        try:
            xs = [float(point[0]) for point in box]
            ys = [float(point[1]) for point in box]
        except (TypeError, ValueError, IndexError):
            continue
        items.append(
            {
                "text": text,
                "x1": min(xs),
                "x2": max(xs),
                "y1": min(ys),
                "y2": max(ys),
                "xc": sum(xs) / len(xs),
                "yc": sum(ys) / len(ys),
            }
        )
    return items


def layout_lines_from_items(items, image_width, image_height):
    if not items:
        return []

    sorted_items = sorted(items, key=lambda item: (item["yc"], item["x1"]))
    line_threshold = max(image_height * 0.008, 12)
    lines = []
    for item in sorted_items:
        current = None
        for line in reversed(lines[-4:]):
            if abs(line["yc"] - item["yc"]) <= line_threshold:
                current = line
                break
        if current is None:
            current = {"items": [], "yc": item["yc"]}
            lines.append(current)
        current["items"].append(item)
        current["yc"] = sum(part["yc"] for part in current["items"]) / len(current["items"])

    output = []
    for line in lines:
        parts = sorted(line["items"], key=lambda item: item["x1"])
        text = " ".join(part["text"] for part in parts)
        output.append(
            {
                "text": text,
                "compact": compact_text(text),
                "x1": min(part["x1"] for part in parts) / image_width,
                "x2": max(part["x2"] for part in parts) / image_width,
                "y1": min(part["y1"] for part in parts) / image_height,
                "y2": max(part["y2"] for part in parts) / image_height,
                "yc": sum(part["yc"] for part in parts) / len(parts) / image_height,
            }
        )
    return sorted(output, key=lambda line: (line["yc"], line["x1"]))


def layout_heading_match(field, compact_line):
    if field == "NumeroCertificado":
        return bool(re.search(r"(N|NO|NRO|NUMERO).{0,8}(ARG|BOL|BRA|CHL|PER|PRY|URY|VEN|COL|ESP|ITA)", compact_line)) or bool(
            re.search(r"(ARG|BOL|BRA|CHL|PER|PRY|URY|VEN|COL|ESP|ITA)\d{2}0*\d{5,8}", compact_line)
        )
    if field == "Diagnostico":
        return "DIAGNOSTICO" in compact_line and "FUNCIONAL" not in compact_line
    if field == "FuncionesCorporales":
        return "FUNCIONESCORPORALES" in compact_line
    if field == "EstructurasCorporales":
        return "ESTRUCTURASCORPORALES" in compact_line
    if field == "ActividadParticipacion":
        return "ACTIVIDAD" in compact_line and "PARTICIP" in compact_line
    if field == "FactoresAmbientales":
        return "FACTORESAMBIENTALES" in compact_line
    if field == "OrientacionPrestacional":
        return "ORIENTACIONPRESTACIONAL" in compact_line
    if field == "Acompaniante":
        return "ACOMPANANTE" in compact_line
    return False


def layout_sections_text(engine, image):
    items = rapidocr_items_with_boxes(engine, image)
    return layout_sections_text_from_items(items, image.width, image.height)


def layout_sections_text_from_items(items, image_width, image_height):
    lines = layout_lines_from_items(items, image_width, image_height)
    if not lines:
        return ""

    headings = []
    for index, line in enumerate(lines):
        for field in LAYOUT_SECTION_FIELDS:
            if layout_heading_match(field, line["compact"]):
                headings.append({"field": field, "index": index, "yc": line["yc"]})
                break

    if not headings:
        return ""

    headings = sorted(headings, key=lambda item: item["index"])
    parts = []
    for pos, heading in enumerate(headings):
        field = heading["field"]
        start_index = heading["index"]
        end_index = headings[pos + 1]["index"] if pos + 1 < len(headings) else len(lines)
        if field == "NumeroCertificado":
            end_index = min(end_index, start_index + 3)
        section_lines = lines[start_index:end_index]
        section_text = "\n".join(line["text"] for line in section_lines).strip()
        if section_text:
            parts.extend(
                [
                    f"[OCR_LAYOUT_SECCION_{field.upper()}]",
                    f"{FIELD_LABELS.get(field, field)}:",
                    section_text,
                ]
            )
    return "\n".join(parts)


def rapidocr_text_with_layout(engine, image):
    items = rapidocr_items_with_boxes(engine, image)
    plain_text = "\n".join(item["text"] for item in items)
    section_text = layout_sections_text_from_items(items, image.width, image.height)
    return "\n".join(part for part in [plain_text, section_text] if part)


def crop_by_ratio(image, ratios):
    width, height = image.size
    left, top, right, bottom = ratios
    return image.crop(
        (
            int(width * left),
            int(height * top),
            int(width * right),
            int(height * bottom),
        )
    )


def ocr_pdf_rapidocr_layout(pdf_path, dpi, max_pages=None):
    try:
        RapidOCR = _PreloadedRapidOCR
        if RapidOCR is None:
            from rapidocr_onnxruntime import RapidOCR
    except ImportError as exc:
        raise RuntimeError(
            f"Falta RapidOCR o una dependencia binaria no cargo correctamente: {exc}"
        ) from exc

    engine = RapidOCR()
    crop_zones = {
        "zona_certificado": (0.39, 0.025, 0.99, 0.125),
        "zona_datos_personales": (0.04, 0.07, 0.99, 0.215),
        "zona_diagnostico_principal": (0.04, 0.135, 0.99, 0.285),
        "zona_diagnostico_cif": (0.04, 0.17, 0.99, 0.44),
        "zona_funciones": (0.05, 0.215, 0.91, 0.315),
        "zona_estructuras": (0.05, 0.255, 0.91, 0.365),
        "zona_actividad_factores": (0.04, 0.30, 0.99, 0.585),
        "zona_actividad_extendida": (0.03, 0.335, 0.99, 0.515),
        "zona_actividad_derecha": (0.45, 0.325, 0.99, 0.515),
        "zona_actividad_inferior": (0.03, 0.405, 0.99, 0.535),
        "zona_factores_extendida": (0.03, 0.425, 0.99, 0.625),
        "zona_factores_derecha": (0.45, 0.425, 0.99, 0.625),
    }

    page_texts = []
    for page_number, image in render_pdf_pages(pdf_path, dpi=dpi, max_pages=max_pages):
        processed = otsu_image(image)
        enhanced = enhanced_ocr_image(image)
        parts = [
            "[OCR_ORIGINAL]",
            rapidocr_text(engine, image),
            "[OCR_ENHANCED]",
            rapidocr_text(engine, enhanced),
            "[OCR_OTSU]",
            rapidocr_text(engine, processed),
        ]

        for zone_name, ratios in crop_zones.items():
            crop = crop_by_ratio(processed, ratios)
            crop = crop.resize((crop.width * 3, crop.height * 3))
            parts.extend([f"[OCR_RECORTE_{zone_name.upper()}]", rapidocr_text(engine, crop)])
            enhanced_crop = crop_by_ratio(enhanced, ratios)
            enhanced_crop = enhanced_crop.resize((enhanced_crop.width * 4, enhanced_crop.height * 4))
            parts.extend([
                f"[OCR_RECORTE_MEJORADO_{zone_name.upper()}]",
                rapidocr_text(engine, enhanced_crop),
            ])
            if zone_name in {
                "zona_certificado",
                "zona_diagnostico_principal",
                "zona_actividad_derecha",
                "zona_actividad_inferior",
                "zona_factores_derecha",
            }:
                sharp_crop = enhanced_crop.filter(ImageFilter.SHARPEN)
                sharp_crop = ImageEnhance.Contrast(sharp_crop).enhance(1.7)
                parts.extend([
                    f"[OCR_RECORTE_NITIDO_{zone_name.upper()}]",
                    rapidocr_text(engine, sharp_crop),
                ])

        page_texts.append({"pagina": page_number, "texto": "\n".join(parts)})
    return page_texts


def ocr_pdf_rapidocr_layout_fast(pdf_path, dpi, max_pages=None):
    try:
        RapidOCR = _PreloadedRapidOCR
        if RapidOCR is None:
            from rapidocr_onnxruntime import RapidOCR
    except ImportError as exc:
        raise RuntimeError(
            f"Falta RapidOCR o una dependencia binaria no cargo correctamente: {exc}"
        ) from exc

    engine = RapidOCR()
    crop_zones = {
        "zona_certificado": (0.39, 0.025, 0.99, 0.125),
        "zona_datos_personales": (0.04, 0.07, 0.99, 0.215),
        "zona_diagnostico_principal": (0.04, 0.135, 0.99, 0.285),
        "zona_cif_general": (0.04, 0.20, 0.99, 0.62),
    }

    page_texts = []
    for page_number, image in render_pdf_pages(pdf_path, dpi=dpi, max_pages=max_pages):
        enhanced = enhanced_ocr_image(image)
        parts = [
            "[OCR_FAST_ENHANCED]",
            rapidocr_text_with_layout(engine, enhanced),
        ]
        for zone_name, ratios in crop_zones.items():
            enhanced_crop = crop_by_ratio(enhanced, ratios)
            enhanced_crop = enhanced_crop.resize((enhanced_crop.width * 3, enhanced_crop.height * 3))
            parts.extend([
                f"[OCR_FAST_RECORTE_{zone_name.upper()}]",
                rapidocr_text(engine, enhanced_crop),
            ])
        page_texts.append({"pagina": page_number, "texto": "\n".join(parts)})
    return page_texts


FIELD_RETRY_ZONES = {
    "NumeroCertificado": {
        "zona_certificado": (0.39, 0.025, 0.99, 0.125),
    },
    "ApellidoNombre": {
        "zona_datos_personales": (0.04, 0.07, 0.99, 0.215),
    },
    "DNI": {
        "zona_datos_personales": (0.04, 0.07, 0.99, 0.215),
    },
    "CUIL": {
        "zona_datos_personales": (0.04, 0.07, 0.99, 0.215),
    },
    "Sexo": {
        "zona_datos_personales": (0.04, 0.07, 0.99, 0.215),
    },
    "FechaNacimiento": {
        "zona_datos_personales": (0.04, 0.07, 0.99, 0.215),
    },
    "Diagnostico": {
        "zona_diagnostico_texto": (0.11, 0.155, 0.98, 0.235),
        "zona_diagnostico_principal": (0.04, 0.135, 0.99, 0.285),
        "zona_diagnostico_cif": (0.04, 0.17, 0.99, 0.44),
    },
    "FuncionesCorporales": {
        "zona_funciones": (0.05, 0.215, 0.91, 0.315),
        "zona_diagnostico_cif": (0.04, 0.17, 0.99, 0.44),
    },
    "EstructurasCorporales": {
        "zona_estructuras": (0.05, 0.255, 0.91, 0.365),
        "zona_diagnostico_cif": (0.04, 0.17, 0.99, 0.44),
    },
    "ActividadParticipacion": {
        "zona_actividad_extendida": (0.03, 0.335, 0.99, 0.515),
        "zona_actividad_derecha": (0.45, 0.325, 0.99, 0.515),
        "zona_actividad_inferior": (0.03, 0.405, 0.99, 0.535),
    },
    "FactoresAmbientales": {
        "zona_factores_extendida": (0.03, 0.425, 0.99, 0.625),
        "zona_factores_derecha": (0.45, 0.425, 0.99, 0.625),
    },
    "OrientacionPrestacional": {
        "zona_orientacion": (0.04, 0.54, 0.99, 0.72),
    },
    "FechaEmision": {
        "zona_fechas": (0.04, 0.70, 0.99, 0.92),
    },
    "FechaVencimiento": {
        "zona_fechas": (0.04, 0.70, 0.99, 0.92),
    },
    "Acompaniante": {
        "zona_acompaniante": (0.04, 0.54, 0.99, 0.78),
    },
}


def retry_zones_for_fields(fields):
    zones = {}
    for field in sorted(set(fields)):
        zones.update(FIELD_RETRY_ZONES.get(field, {}))
    return zones


def targeted_engine_name(fields):
    safe_fields = "_".join(sorted(set(str(field) for field in fields)))
    safe_fields = re.sub(r"[^A-Za-z0-9_]+", "_", safe_fields)
    if len(safe_fields) > 120:
        safe_fields = str(abs(hash(safe_fields)))
    return f"rapidocr_targeted_{safe_fields or 'sin_campos'}"


def ocr_pdf_rapidocr_targeted(pdf_path, dpi, fields, max_pages=None):
    try:
        RapidOCR = _PreloadedRapidOCR
        if RapidOCR is None:
            from rapidocr_onnxruntime import RapidOCR
    except ImportError as exc:
        raise RuntimeError(
            f"Falta RapidOCR o una dependencia binaria no cargo correctamente: {exc}"
        ) from exc

    zones = retry_zones_for_fields(fields)
    if not zones:
        return []

    engine = RapidOCR()
    if any(field in CIF_FIELDS or field == "Diagnostico" for field in fields):
        dpi = max(dpi, 500)
    page_texts = []
    for page_number, image in render_pdf_pages(pdf_path, dpi=dpi, max_pages=max_pages):
        enhanced = enhanced_ocr_image(image)
        parts = []
        for zone_name, ratios in zones.items():
            enhanced_crop = crop_by_ratio(enhanced, ratios)
            enhanced_crop = enhanced_crop.resize((enhanced_crop.width * 4, enhanced_crop.height * 4))
            parts.extend([
                f"[OCR_TARGET_RECORTE_{zone_name.upper()}]",
                rapidocr_text_with_layout(engine, enhanced_crop),
            ])
            sharp_crop = enhanced_crop.filter(ImageFilter.SHARPEN)
            sharp_crop = ImageEnhance.Contrast(sharp_crop).enhance(1.7)
            parts.extend([
                f"[OCR_TARGET_NITIDO_{zone_name.upper()}]",
                rapidocr_text(engine, sharp_crop),
            ])
        page_texts.append({"pagina": page_number, "texto": "\n".join(parts)})
    return page_texts


def ocr_pdf_targeted_cached(pdf_path, dpi, fields, lang, poppler_path=None, tesseract_cmd=None, max_pages=None, cache_dir=None):
    ocr_engine = targeted_engine_name(fields)
    if cache_dir is None:
        return ocr_pdf_rapidocr_targeted(pdf_path, dpi=dpi, fields=fields, max_pages=max_pages)

    cache_dir.mkdir(parents=True, exist_ok=True)
    key = cache_key(pdf_path, dpi, max_pages, ocr_engine)
    sqlite_page_texts = read_sqlite_ocr_cache(cache_dir, key, pdf_path)
    if sqlite_page_texts is not None:
        return sqlite_page_texts

    cache_path = cache_dir / key
    if cache_path.exists():
        page_texts = json.loads(cache_path.read_text(encoding="utf-8"))
        write_sqlite_ocr_cache(cache_dir, key, pdf_path, dpi, max_pages, ocr_engine, page_texts)
        return page_texts

    page_texts = ocr_pdf_rapidocr_targeted(pdf_path, dpi=dpi, fields=fields, max_pages=max_pages)
    cache_path.write_text(json.dumps(page_texts, ensure_ascii=False, indent=2), encoding="utf-8")
    write_sqlite_ocr_cache(cache_dir, key, pdf_path, dpi, max_pages, ocr_engine, page_texts)
    return page_texts


def ocr_pdf(
    pdf_path,
    dpi,
    lang,
    poppler_path=None,
    tesseract_cmd=None,
    max_pages=None,
    ocr_engine="rapidocr",
):
    if ocr_engine == "rapidocr":
        return ocr_pdf_rapidocr(pdf_path=pdf_path, dpi=dpi, max_pages=max_pages)

    if ocr_engine == "rapidocr_layout":
        return ocr_pdf_rapidocr_layout(pdf_path=pdf_path, dpi=dpi, max_pages=max_pages)

    if ocr_engine == "rapidocr_layout_fast":
        return ocr_pdf_rapidocr_layout_fast(pdf_path=pdf_path, dpi=dpi, max_pages=max_pages)

    if ocr_engine == "tesseract_pymupdf":
        return ocr_pdf_tesseract(
            pdf_path=pdf_path,
            dpi=dpi,
            lang=lang,
            tesseract_cmd=tesseract_cmd,
            max_pages=max_pages,
        )

    try:
        from pdf2image import convert_from_path
        import pytesseract
    except ImportError as exc:
        raise RuntimeError(
            "Faltan dependencias. Instale con: pip install pdf2image pytesseract pillow"
        ) from exc

    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd

    pages = convert_from_path(
        str(pdf_path),
        dpi=dpi,
        poppler_path=poppler_path,
        first_page=1,
        last_page=max_pages,
    )

    page_texts = []
    for page_number, page in enumerate(pages, start=1):
        text = pytesseract.image_to_string(page, lang=lang)
        page_texts.append({"pagina": page_number, "texto": text})
    return page_texts


def cache_key(pdf_path, dpi, max_pages, ocr_engine):
    safe_name = re.sub(r"[^A-Za-z0-9_.-]+", "_", pdf_path.name)
    pages = "all" if max_pages is None else str(max_pages)
    return f"{safe_name}.dpi{dpi}.pages{pages}.{ocr_engine}.{OCR_CACHE_VERSION}.json"


def sqlite_cache_path(cache_dir):
    return cache_dir / "ocr_cache.sqlite"


def init_sqlite_cache(cache_dir):
    db_path = sqlite_cache_path(cache_dir)
    with sqlite3.connect(db_path) as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS ocr_cache (
                cache_key TEXT PRIMARY KEY,
                pdf_name TEXT NOT NULL,
                pdf_path TEXT NOT NULL,
                pdf_size INTEGER NOT NULL,
                pdf_mtime REAL NOT NULL,
                dpi INTEGER NOT NULL,
                max_pages TEXT NOT NULL,
                ocr_engine TEXT NOT NULL,
                cache_version TEXT NOT NULL,
                created_at TEXT NOT NULL,
                page_texts_json TEXT NOT NULL
            )
            """
        )
        conn.execute(
            "CREATE INDEX IF NOT EXISTS idx_ocr_cache_pdf_name ON ocr_cache(pdf_name)"
        )
    return db_path


def pdf_metadata(pdf_path):
    stat = pdf_path.stat()
    return stat.st_size, stat.st_mtime


def read_sqlite_ocr_cache(cache_dir, key, pdf_path):
    db_path = init_sqlite_cache(cache_dir)
    pdf_size, pdf_mtime = pdf_metadata(pdf_path)
    with sqlite3.connect(db_path) as conn:
        row = conn.execute(
            """
            SELECT page_texts_json
            FROM ocr_cache
            WHERE cache_key = ?
              AND pdf_size = ?
              AND pdf_mtime = ?
              AND cache_version = ?
            """,
            (key, pdf_size, pdf_mtime, OCR_CACHE_VERSION),
        ).fetchone()
    if not row:
        return None
    return json.loads(row[0])


def write_sqlite_ocr_cache(cache_dir, key, pdf_path, dpi, max_pages, ocr_engine, page_texts):
    db_path = init_sqlite_cache(cache_dir)
    pdf_size, pdf_mtime = pdf_metadata(pdf_path)
    pages = "all" if max_pages is None else str(max_pages)
    with sqlite3.connect(db_path) as conn:
        conn.execute(
            """
            INSERT OR REPLACE INTO ocr_cache (
                cache_key,
                pdf_name,
                pdf_path,
                pdf_size,
                pdf_mtime,
                dpi,
                max_pages,
                ocr_engine,
                cache_version,
                created_at,
                page_texts_json
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                key,
                pdf_path.name,
                str(pdf_path),
                pdf_size,
                pdf_mtime,
                dpi,
                pages,
                ocr_engine,
                OCR_CACHE_VERSION,
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                json.dumps(page_texts, ensure_ascii=False),
            ),
        )


def ocr_cache_exists(cache_dir, pdf_path, dpi, max_pages, ocr_engine):
    key = cache_key(pdf_path, dpi, max_pages, ocr_engine)
    db_path = init_sqlite_cache(cache_dir)
    pdf_size, pdf_mtime = pdf_metadata(pdf_path)
    with sqlite3.connect(db_path) as conn:
        row = conn.execute(
            """
            SELECT 1
            FROM ocr_cache
            WHERE cache_key = ?
              AND pdf_size = ?
              AND pdf_mtime = ?
              AND cache_version = ?
            LIMIT 1
            """,
            (key, pdf_size, pdf_mtime, OCR_CACHE_VERSION),
        ).fetchone()
    if row:
        return True
    return (cache_dir / key).exists()


def ocr_pdf_cached(
    pdf_path,
    dpi,
    lang,
    poppler_path=None,
    tesseract_cmd=None,
    max_pages=None,
    ocr_engine="rapidocr",
    cache_dir=None,
):
    if cache_dir is None:
        return ocr_pdf(
            pdf_path=pdf_path,
            dpi=dpi,
            lang=lang,
            poppler_path=poppler_path,
            tesseract_cmd=tesseract_cmd,
            max_pages=max_pages,
            ocr_engine=ocr_engine,
    )

    cache_dir.mkdir(parents=True, exist_ok=True)
    key = cache_key(pdf_path, dpi, max_pages, ocr_engine)
    sqlite_page_texts = read_sqlite_ocr_cache(cache_dir, key, pdf_path)
    if sqlite_page_texts is not None:
        return sqlite_page_texts

    cache_path = cache_dir / key
    if cache_path.exists():
        page_texts = json.loads(cache_path.read_text(encoding="utf-8"))
        write_sqlite_ocr_cache(
            cache_dir,
            key,
            pdf_path,
            dpi,
            max_pages,
            ocr_engine,
            page_texts,
        )
        return page_texts

    page_texts = ocr_pdf(
        pdf_path=pdf_path,
        dpi=dpi,
        lang=lang,
        poppler_path=poppler_path,
        tesseract_cmd=tesseract_cmd,
        max_pages=max_pages,
        ocr_engine=ocr_engine,
    )
    cache_path.write_text(
        json.dumps(page_texts, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    write_sqlite_ocr_cache(
        cache_dir,
        key,
        pdf_path,
        dpi,
        max_pages,
        ocr_engine,
        page_texts,
    )
    return page_texts


def extract_matches(page_texts, pattern):
    regex = re.compile(pattern, flags=re.IGNORECASE)
    rows = []
    for page in page_texts:
        for match in regex.findall(page["texto"]):
            if isinstance(match, tuple):
                match = next((part for part in match if part), "")
            value = normalize_identifier(match)
            if value:
                rows.append({"identificador": value, "pagina_pdf": page["pagina"]})
    return rows


def first_identifier_from_name(pdf_path, pattern):
    regex = re.compile(pattern, flags=re.IGNORECASE)
    matches = regex.findall(pdf_path.stem)
    if not matches:
        return ""
    match = matches[0]
    if isinstance(match, tuple):
        match = next((part for part in match if part), "")
    return normalize_identifier(match)


def pdf_dni_sort_key(pdf_path, pattern):
    identifier = first_identifier_from_name(pdf_path, pattern)
    if identifier:
        return (0, int(identifier), pdf_path.stem.lower())
    return (1, pdf_path.stem.lower())


def sorted_pdf_paths_by_dni(pdf_dir, pattern):
    return sorted(pdf_dir.glob("*.pdf"), key=lambda path: pdf_dni_sort_key(path, pattern))


def build_pdf_validation(pdf_paths, df_team_ids, df_ocr, pattern):
    team_ids = set(df_team_ids["identificador"])
    rows = []
    for pdf_path in pdf_paths:
        expected_id = first_identifier_from_name(pdf_path, pattern)
        matches = df_ocr[
            (df_ocr["pdf"] == pdf_path.name) & (df_ocr["identificador"] == expected_id)
        ]
        pages = sorted(
            {
                int(value)
                for value in matches["pagina_pdf"].dropna().tolist()
                if str(value).strip() != ""
            }
        )

        if not expected_id:
            status = "SIN_ID_EN_NOMBRE_PDF"
        elif expected_id not in team_ids:
            status = "PDF_NO_EXISTE_EN_ARCHIVO"
        elif not pages and not matches.empty:
            status = "OK_ID_DESDE_NOMBRE"
        elif pages:
            status = "OK_ID_EN_OCR"
        else:
            status = "ID_DEL_PDF_NO_LEIDO_POR_OCR"

        rows.append(
            {
                "pdf": pdf_path.name,
                "identificador_esperado": expected_id,
                "existe_en_archivo": "SI" if expected_id in team_ids else "NO",
                "identificador_leido_por_ocr": "SI" if not matches.empty else "NO",
                "paginas_ocr": ", ".join(map(str, pages)),
                "estado_pdf": status,
            }
        )
    return pd.DataFrame(rows)


def certificate_column_name(df):
    for column in ["Origen - NumeroCertificado", "Origen - Numero Certificado"]:
        if column in df.columns:
            return column
    return None


def certificate_matches_ocr(certificate_value, ocr_text):
    certificate_text = compact_text(certificate_value)
    if not certificate_text:
        return False
    return certificate_text in compact_text(ocr_text)


def audit_key_from_parts(identifier, certificate_value="", pdf_name=""):
    identifier = normalize_identifier(identifier)
    certificate = str(certificate_value or "").strip()
    if certificate:
        return f"{identifier} | {certificate}"
    if pdf_name:
        return f"{identifier} | {pdf_name}"
    return identifier


def build_team_rows_for_pdf_selection(pdf_paths, df_team, raw_text_rows, pattern):
    raw_by_pdf = {}
    for row in raw_text_rows:
        raw_by_pdf.setdefault(row["pdf"], "")
        raw_by_pdf[row["pdf"]] += "\n" + str(row.get("texto_ocr", ""))

    cert_column = certificate_column_name(df_team)
    selected_rows = []
    for pdf_path in pdf_paths:
        expected_id = first_identifier_from_name(pdf_path, pattern)
        if not expected_id:
            continue

        matches = df_team[df_team["identificador"] == expected_id]
        if matches.empty:
            continue

        ocr_text = raw_by_pdf.get(pdf_path.name, "")
        selected = None
        status = "DNI_UNICO"
        matched_certificates = []
        if len(matches) == 1:
            selected = matches.iloc[0].copy()
        else:
            status = "DNI_DUPLICADO_NO_RESUELTO_POR_CERTIFICADO"
            if cert_column:
                certificate_matches = []
                for row_index, candidate in matches.iterrows():
                    certificate = candidate.get(cert_column, "")
                    if certificate_matches_ocr(certificate, ocr_text):
                        certificate_matches.append((row_index, candidate, str(certificate).strip()))
                if len(certificate_matches) == 1:
                    _, selected, certificate = certificate_matches[0]
                    selected = selected.copy()
                    matched_certificates = [certificate]
                    status = "DNI_DUPLICADO_RESUELTO_POR_CERTIFICADO"
                elif len(certificate_matches) > 1:
                    selected = matches.iloc[0].copy()
                    matched_certificates = [certificate for _, _, certificate in certificate_matches]
                    status = "DNI_DUPLICADO_MULTIPLES_CERTIFICADOS_EN_OCR"

            if selected is None:
                selected = matches.iloc[0].copy()

        selected["PDF auditado"] = pdf_path.name
        selected["Fila Excel seleccionada"] = int(selected.name) + 2
        selected["Cantidad filas mismo DNI"] = len(matches)
        selected["Estado seleccion registro"] = status
        selected["Certificado usado para resolver DNI"] = " | ".join(matched_certificates)
        selected_certificate = selected.get(cert_column, "") if cert_column else ""
        selected["clave_auditoria"] = audit_key_from_parts(
            expected_id,
            selected_certificate,
            pdf_path.name,
        )
        selected_rows.append(selected)

    if not selected_rows:
        return pd.DataFrame()
    return pd.DataFrame(selected_rows).drop_duplicates(subset=["clave_auditoria"], keep="last")


def build_field_comparison(pdf_paths, df_team_ids, raw_text_rows, pattern):
    raw_by_pdf = {}
    for row in raw_text_rows:
        raw_by_pdf.setdefault(row["pdf"], "")
        raw_by_pdf[row["pdf"]] += "\n" + str(row.get("texto_ocr", ""))

    team_by_pdf = {
        str(row.get("PDF auditado", "")): row
        for _, row in df_team_ids.iterrows()
        if str(row.get("PDF auditado", "")).strip()
    }
    team_by_id = {}
    for _, row in df_team_ids.iterrows():
        identifier = str(row["identificador"]).strip()
        if identifier and identifier not in team_by_id:
            team_by_id[identifier] = row

    rows = []
    cif_detail_rows = []
    cert_column = certificate_column_name(df_team_ids)
    for pdf_path in pdf_paths:
        expected_id = first_identifier_from_name(pdf_path, pattern)
        team_row = team_by_pdf.get(pdf_path.name)
        if team_row is None:
            team_row = team_by_id.get(expected_id)
        ocr_text = raw_by_pdf.get(pdf_path.name, "")

        if team_row is None:
            rows.append(
                {
                    "pdf": pdf_path.name,
                    "identificador": expected_id,
                    "campo": "_registro",
                    "valor_origen": "",
                    "estado": "PDF_NO_EXISTE_EN_ARCHIVO",
                    "score": 0,
                    "criterio": "",
                }
            )
            continue

        certificate_value = team_row.get(cert_column, "") if cert_column else ""
        audit_key = str(team_row.get("clave_auditoria", "")) or audit_key_from_parts(
            expected_id,
            certificate_value,
            pdf_path.name,
        )
        selection_status = str(team_row.get("Estado seleccion registro", ""))
        if selection_status in {
            "DNI_DUPLICADO_NO_RESUELTO_POR_CERTIFICADO",
            "DNI_DUPLICADO_MULTIPLES_CERTIFICADOS_EN_OCR",
        }:
            rows.append(
                {
                    "pdf": pdf_path.name,
                    "identificador": expected_id,
                    "clave_auditoria": audit_key,
                    "campo": "_registro",
                    "columna_origen": "Origen - DNI",
                    "valor_origen": expected_id,
                    "numero_certificado_origen": certificate_value,
                    "estado": selection_status,
                    "score": 0,
                    "criterio": "DNI duplicado: no se pudo elegir una unica fila por NumeroCertificado",
                    "detalle_revision": (
                        "No se valida automaticamente: el DNI esta duplicado en el archivo "
                        "y el OCR no permitio identificar un unico NumeroCertificado."
                    ),
                    "detalle_pdf_no_en_archivo": "",
                }
            )
            continue

        for field in FIELDS_TO_VALIDATE:
            origin_column = f"Origen - {field}"
            expected = team_row.get(origin_column, "")
            status, score, criterion = field_score(field, expected, ocr_text)
            detail_row = cif_detail_row(
                pdf_path.name,
                expected_id,
                audit_key,
                field,
                expected,
                ocr_text,
            )
            if detail_row is not None:
                detail_row["estado"] = status
                detail_row["detalle_revision"] = "" if status == "OK" else missing_detail(field, expected, ocr_text)
                cif_detail_rows.append(detail_row)
            rows.append(
                {
                    "pdf": pdf_path.name,
                    "identificador": expected_id,
                    "clave_auditoria": audit_key,
                    "campo": field,
                    "columna_origen": origin_column,
                    "valor_origen": expected,
                    "numero_certificado_origen": certificate_value,
                    "estado": status,
                    "score": score,
                    "criterio": criterion,
                    "detalle_revision": "" if status == "OK" else missing_detail(field, expected, ocr_text),
                    "detalle_pdf_no_en_archivo": extra_cif_detail(field, expected, ocr_text),
                }
            )

    return pd.DataFrame(rows), pd.DataFrame(cif_detail_rows)


def build_summary_by_pdf(df_field_comparison):
    if df_field_comparison.empty:
        return pd.DataFrame()

    index_columns = ["pdf", "identificador"]
    if "clave_auditoria" in df_field_comparison.columns:
        index_columns.append("clave_auditoria")

    summary = (
        df_field_comparison.pivot_table(
            index=index_columns,
            columns="estado",
            values="campo",
            aggfunc="count",
            fill_value=0,
        )
        .reset_index()
        .rename_axis(None, axis=1)
    )

    for column in ["OK", "PARCIAL", "NO_ENCONTRADO", "SIN_DATO_ORIGEN"]:
        if column not in summary.columns:
            summary[column] = 0

    summary["total_campos"] = summary[["OK", "PARCIAL", "NO_ENCONTRADO", "SIN_DATO_ORIGEN"]].sum(axis=1)
    summary["estado_general"] = summary.apply(
        lambda row: "OK"
        if row["NO_ENCONTRADO"] == 0 and row["PARCIAL"] == 0
        else ("REVISAR" if row["OK"] > 0 else "CRITICO"),
        axis=1,
    )
    return summary[
        [
            "pdf",
            "identificador",
            *[column for column in ["clave_auditoria"] if column in summary.columns],
            "estado_general",
            "OK",
            "PARCIAL",
            "NO_ENCONTRADO",
            "SIN_DATO_ORIGEN",
            "total_campos",
        ]
    ]


def build_simple_review(df_field_comparison, df_team_ids):
    if df_field_comparison.empty:
        return pd.DataFrame()

    key_column = "clave_auditoria" if "clave_auditoria" in df_team_ids.columns else "identificador"
    order = (
        df_team_ids.reset_index()
        .rename(columns={"index": "orden_archivo"})
        [["orden_archivo", key_column]]
    )

    grouped_rows = []
    group_columns = ["pdf", "identificador"]
    if "clave_auditoria" in df_field_comparison.columns:
        group_columns.append("clave_auditoria")
    for group_key, group in df_field_comparison.groupby(group_columns, sort=False):
        if len(group_columns) == 3:
            pdf, identifier, audit_key = group_key
        else:
            pdf, identifier = group_key
            audit_key = str(identifier)
        alerts = group[group["estado"] != "OK"].copy()
        ok_count = int((group["estado"] == "OK").sum())
        cif_alerts = alerts[
            alerts["campo"].isin(
                [
                    "FuncionesCorporales",
                    "EstructurasCorporales",
                    "ActividadParticipacion",
                    "FactoresAmbientales",
                ]
            )
        ]
        other_alerts = alerts[
            ~alerts["campo"].isin(
                [
                    "FuncionesCorporales",
                    "EstructurasCorporales",
                    "ActividadParticipacion",
                    "FactoresAmbientales",
                ]
            )
        ]

        if alerts.empty:
            result = "OK"
            cif_summary = ""
            other_summary = ""
            summary = "Todos los campos comparados coinciden con el PDF."
        else:
            only_ocr_uncertain = all(is_ocr_uncertain_alert(row) for row in alerts.itertuples(index=False))
            result = "NO_VALIDABLE_OCR" if only_ocr_uncertain else "REVISAR"
            cif_summary = " | ".join(
                f"{row.campo}: {row.detalle_revision}"
                for row in cif_alerts.itertuples(index=False)
            )
            other_summary = " | ".join(
                f"{row.campo}: {row.detalle_revision}"
                for row in other_alerts.itertuples(index=False)
            )
            summary = "Revisar campos indicados."

        grouped_rows.append(
            {
                "identificador": str(identifier),
                "clave_auditoria": str(audit_key),
                "pdf": pdf,
                "resultado": result,
                "campos_ok": ok_count,
                "campos_a_revisar": int(len(alerts)),
            "codigos_cif_del_archivo_no_en_pdf": cif_summary,
                "otros_datos_a_revisar": other_summary,
                "resumen_revision": summary,
            }
        )

    df_simple = pd.DataFrame(grouped_rows)
    df_simple = order.merge(df_simple, on=key_column, how="inner")
    df_simple = df_simple.sort_values("orden_archivo")
    return df_simple[
        [
            *[column for column in ["clave_auditoria"] if column in df_simple.columns],
            "identificador",
            "pdf",
            "resultado",
            "campos_ok",
            "campos_a_revisar",
            "codigos_cif_del_archivo_no_en_pdf",
            "otros_datos_a_revisar",
            "resumen_revision",
        ]
    ]


def normalize_manual_review(value):
    text = normalize_text(value)
    if not text:
        return ""
    if text in {"OK", "SI", "SÍ"}:
        return "OK"
    if text in {"NO OK", "NOOK", "NO_OK", "NOK"}:
        return "NO_OK"
    return text


def manual_expected_result(value):
    normalized = normalize_manual_review(value)
    if normalized == "OK":
        return "OK"
    if normalized == "NO_OK":
        return "REVISAR"
    return ""


def compare_manual_vs_ocr(df_simple_review, df_team_ids, reviewed_column, comments_column):
    if not reviewed_column or reviewed_column not in df_team_ids.columns:
        return pd.DataFrame()

    cols = ["identificador", reviewed_column]
    if comments_column and comments_column in df_team_ids.columns:
        cols.append(comments_column)

    manual = df_team_ids[cols].copy()
    manual["revision_manual"] = manual[reviewed_column].map(normalize_manual_review)
    manual = manual[manual["revision_manual"] != ""]
    if comments_column and comments_column in manual.columns:
        manual["observacion_auditoria"] = manual[comments_column].fillna("")
    else:
        manual["observacion_auditoria"] = ""

    merged = manual.merge(df_simple_review, on="identificador", how="left")
    merged["resultado_esperado_por_manual"] = merged[reviewed_column].map(manual_expected_result)

    def coincidence(row):
        if pd.isna(row.get("pdf")):
            return "PDF_NO_ENCONTRADO"
        if row.get("resultado") == "NO_VALIDABLE_OCR":
            return "NO_APLICA_OCR_NO_VALIDABLE"
        expected = row.get("resultado_esperado_por_manual")
        if not expected:
            return "REVISION_MANUAL_NO_INTERPRETABLE"
        return "SI" if row.get("resultado") == expected else "NO"

    merged["coincide_manual_vs_ocr"] = merged.apply(coincidence, axis=1)
    output_columns = [
        "identificador",
        "pdf",
        reviewed_column,
        "revision_manual",
        "observacion_auditoria",
        "resultado",
        "coincide_manual_vs_ocr",
        "campos_ok",
        "campos_a_revisar",
        "codigos_cif_del_archivo_no_en_pdf",
        "otros_datos_a_revisar",
        "resumen_revision",
    ]
    return merged[[column for column in output_columns if column in merged.columns]]


def build_review_matrix(df_team_ids, df_field_comparison, df_simple_review, reviewed_column, comments_column):
    if df_team_ids.empty:
        return pd.DataFrame()

    base = df_team_ids.copy()
    if "identificador" not in base.columns:
        return pd.DataFrame()

    key_column = "clave_auditoria" if "clave_auditoria" in base.columns else "identificador"
    details_by_id = {}
    for identifier, group in df_field_comparison.groupby(key_column):
        missing_in_pdf_details = []
        ocr_uncertain_details = []
        missing_in_file_details = []
        for row in group.itertuples(index=False):
            field = row.campo
            label = FIELD_LABELS.get(field, field)
            detail = str(row.detalle_revision or "")
            if row.estado == "OK":
                pass
            elif detail.startswith("No validable por OCR"):
                ocr_uncertain_details.append(f"{label}: {detail}")
            else:
                cleaned_detail = detail.replace(
                    "Codigo del certificado no presente en archivo:",
                    "Codigo del certificado no presente en archivo:",
                )
                if "Codigo del certificado no presente en archivo:" not in cleaned_detail:
                    missing_in_pdf_details.append(f"{label}: {cleaned_detail}")

            extra_detail = str(getattr(row, "detalle_pdf_no_en_archivo", "") or "")
            if extra_detail:
                missing_in_file_details.append(extra_detail)

        details_by_id[str(identifier)] = {
            "faltante_en_pdf": " | ".join(dict.fromkeys(missing_in_pdf_details)),
            "campos_no_validables_ocr": " | ".join(ocr_uncertain_details),
            "faltante_en_archivo_entrada": " | ".join(dict.fromkeys(missing_in_file_details)),
        }

    simple_by_id = {
        str(getattr(row, key_column)): row
        for row in df_simple_review.itertuples(index=False)
    }

    rows = []
    for _, source_row in base.iterrows():
        identifier = str(source_row.get(key_column, source_row["identificador"]))
        detail = details_by_id.get(identifier, {"field_markers": {}})
        simple = simple_by_id.get(identifier)
        result = getattr(simple, "resultado", "") if simple is not None else ""
        manual_value = source_row.get(reviewed_column, "") if reviewed_column in source_row.index else ""
        manual_expected = manual_expected_result(manual_value)
        if not manual_expected or not result:
            manual_match = ""
        elif result == "NO_VALIDABLE_OCR":
            manual_match = "NO APLICA - OCR NO VALIDABLE"
        else:
            manual_match = "SI" if result == manual_expected else "NO"

        row = source_row.to_dict()
        row["Resultado OCR"] = result
        row["Falta en archivo de entrada"] = detail.get("faltante_en_archivo_entrada", "")
        row["No encontrado en PDF"] = detail.get("faltante_en_pdf", "")
        row["Campos no validables por OCR"] = detail.get("campos_no_validables_ocr", "")
        rows.append(row)

    matrix = pd.DataFrame(rows)
    manual_field_columns = set(FIELD_LABELS.keys()) | {
        "Acompañante",
        "Acompaniante",
    }
    original_cols = [
        column
        for column in df_team_ids.columns
        if column not in {"identificador", "clave_auditoria"} and column not in manual_field_columns
    ]
    final_cols = [
        *original_cols,
        "Resultado OCR",
        "Falta en archivo de entrada",
        "No encontrado en PDF",
        "Campos no validables por OCR",
    ]
    return matrix[[column for column in final_cols if column in matrix.columns]]


def build_checkbox_view(df_review_matrix, reviewed_column, comments_column):
    if df_review_matrix.empty:
        return pd.DataFrame()
    return df_review_matrix.copy()


def add_run_metadata(df, run_timestamp, lote, skip_pdfs, max_pdfs):
    output = df.copy()
    if "Origen - DNI" in output.columns:
        dni_values = output["Origen - DNI"].map(normalize_identifier)
    elif "identificador" in output.columns:
        dni_values = output["identificador"].map(normalize_identifier)
    else:
        dni_values = pd.Series([""] * len(output), index=output.index)

    cert_column = certificate_column_name(output)
    if cert_column:
        certificate_values = output[cert_column].fillna("").map(lambda value: str(value).strip())
    elif "numero_certificado_origen" in output.columns:
        certificate_values = output["numero_certificado_origen"].fillna("").map(lambda value: str(value).strip())
    else:
        certificate_values = pd.Series([""] * len(output), index=output.index)

    audit_ids = []
    for dni, certificate in zip(dni_values, certificate_values):
        audit_ids.append(f"{dni} | {certificate}" if certificate else dni)

    metadata = {
        "ID Auditoria": audit_ids,
        "DNI Auditoria": dni_values,
        "Fecha ejecucion": run_timestamp,
        "Lote": lote,
        "Skip": skip_pdfs,
        "Limit": "" if max_pdfs is None else max_pdfs,
    }
    for position, (column, value) in enumerate(metadata.items()):
        output.insert(position, column, value)
    return output


def ensure_user_review_columns(df):
    output = df.copy()
    insert_at = 6 if len(output.columns) >= 6 else len(output.columns)
    for offset, column in enumerate(USER_REVIEW_COLUMNS):
        if column not in output.columns:
            output.insert(insert_at + offset, column, "")
    return output


def is_filled_value(value):
    if pd.isna(value):
        return False
    return str(value).strip() != ""


def preserve_manual_values(previous_df, current_df, dedupe_columns, preserved_columns):
    if previous_df.empty or current_df.empty:
        return current_df

    available_keys = [
        column
        for column in dedupe_columns
        if column in previous_df.columns and column in current_df.columns
    ]
    if not available_keys:
        return current_df

    previous_by_key = {}
    for _, previous_row in previous_df.iterrows():
        key = tuple(str(previous_row.get(column, "")) for column in available_keys)
        previous_by_key[key] = previous_row

    output = current_df.copy()
    for column in preserved_columns:
        if column not in output.columns:
            output[column] = ""
        if column not in previous_df.columns:
            continue
        for index, current_row in output.iterrows():
            key = tuple(str(current_row.get(key_column, "")) for key_column in available_keys)
            previous_row = previous_by_key.get(key)
            if previous_row is not None and is_filled_value(previous_row.get(column, "")):
                output.at[index, column] = previous_row.get(column, "")
    return output


def merge_consolidated_sheet(output_path, sheet_name, current_df, dedupe_columns, preserved_columns=None):
    if not output_path.exists():
        return current_df
    try:
        previous_df = pd.read_excel(output_path, sheet_name=sheet_name, dtype=object)
    except ValueError:
        return current_df

    if preserved_columns:
        current_df = preserve_manual_values(
            previous_df,
            current_df,
            dedupe_columns,
            preserved_columns,
        )

    combined = pd.concat([previous_df, current_df], ignore_index=True)
    available_dedupe_columns = [
        column for column in dedupe_columns if column in combined.columns
    ]
    if available_dedupe_columns:
        combined = combined.drop_duplicates(
            subset=available_dedupe_columns,
            keep="last",
        )
    return combined


def processed_pdf_names_from_output(output_path):
    if not output_path.exists():
        return set()
    try:
        previous_df = pd.read_excel(output_path, sheet_name="Auditoria", dtype=object)
    except ValueError:
        return set()
    if "PDF auditado" not in previous_df.columns:
        return set()
    return {
        str(value).strip()
        for value in previous_df["PDF auditado"].dropna().tolist()
        if str(value).strip()
  }


def normalize_comment(value):
    if pd.isna(value):
        return ""
    return re.sub(r"\s+", " ", str(value).strip())


def processed_pdf_names_to_omit(
    output_path,
    df_team,
    pdf_paths,
    pattern,
    reviewed_column,
    comments_column,
    solo_revisados=False,
):
    processed_names = processed_pdf_names_from_output(output_path)
    return processed_names


def build_duplicate_id_report(
    df_team,
    id_column,
    pdf_paths=None,
    pattern=DEFAULT_ID_PATTERN,
    reviewed_column="REVISADO",
):
    if df_team.empty or "identificador" not in df_team.columns:
        return pd.DataFrame()

    id_counts = df_team["identificador"].value_counts()
    duplicated_ids = set(id_counts[id_counts > 1].index)
    if not duplicated_ids:
        return pd.DataFrame()

    pdfs_by_id = {}
    for pdf_path in pdf_paths or []:
        pdf_id = first_identifier_from_name(pdf_path, pattern)
        if pdf_id:
            pdfs_by_id.setdefault(pdf_id, []).append(pdf_path.name)

    cert_column = "Origen - NumeroCertificado"
    if cert_column not in df_team.columns:
        cert_column = "Origen - Numero Certificado"

    certs_by_id = {}
    if cert_column in df_team.columns:
        for identifier, group in df_team.groupby("identificador", sort=False):
            values = [
                str(value).strip()
                for value in group[cert_column].tolist()
                if str(value).strip() and not pd.isna(value)
            ]
            certs_by_id[identifier] = sorted(set(values))

    reviewed_counts_by_id = {}
    if reviewed_column in df_team.columns:
        reviewed_flags = df_team[reviewed_column].map(normalize_manual_review) != ""
        reviewed_counts_by_id = (
            df_team[reviewed_flags]
            .groupby("identificador", sort=False)
            .size()
            .to_dict()
        )

    def risk_for(identifier):
        pdf_count = len(pdfs_by_id.get(identifier, []))
        cert_count = len(certs_by_id.get(identifier, []))
        if pdf_count > 1 and cert_count > 1:
            return "ALTO: DNI con varios PDFs y varios certificados"
        if cert_count > 1:
            return "MEDIO: DNI con varios certificados en entrada"
        if pdf_count > 1:
            return "MEDIO: DNI con varios PDFs en carpeta"
        return "REVISAR: DNI duplicado en entrada"

    duplicate_rows = df_team[df_team["identificador"].isin(duplicated_ids)].copy()
    duplicate_rows.insert(0, "ID Auditoria", duplicate_rows["identificador"])
    duplicate_rows.insert(1, "Cantidad apariciones", duplicate_rows["identificador"].map(id_counts))
    duplicate_rows.insert(2, "Fila Excel estimada", duplicate_rows.index + 2)
    duplicate_rows.insert(3, "Columna ID usada", id_column)
    duplicate_rows.insert(
        4,
        "Cantidad PDFs encontrados",
        duplicate_rows["identificador"].map(lambda value: len(pdfs_by_id.get(value, []))),
    )
    duplicate_rows.insert(
        5,
        "PDFs encontrados por nombre",
        duplicate_rows["identificador"].map(lambda value: " | ".join(pdfs_by_id.get(value, []))),
    )
    duplicate_rows.insert(
        6,
        "Cantidad certificados distintos",
        duplicate_rows["identificador"].map(lambda value: len(certs_by_id.get(value, []))),
    )
    duplicate_rows.insert(
        7,
        "Numeros certificado en entrada",
        duplicate_rows["identificador"].map(lambda value: " | ".join(certs_by_id.get(value, []))),
    )
    duplicate_rows.insert(
        8,
        "Filas con REVISADO informado",
        duplicate_rows["identificador"].map(lambda value: reviewed_counts_by_id.get(value, 0)),
    )
    duplicate_rows.insert(
        9,
        "Riesgo cruce por DNI",
        duplicate_rows["identificador"].map(risk_for),
    )
    return duplicate_rows.sort_values(["ID Auditoria", "Fila Excel estimada"])


def add_duplicate_id_flags(df_team_ids, df_duplicate_ids):
    output = df_team_ids.copy()
    output["DNI duplicado en entrada"] = ""
    output["Riesgo cruce por DNI"] = ""
    output["PDFs encontrados por DNI"] = ""
    output["Certificados del DNI en entrada"] = ""

    if df_duplicate_ids.empty:
        return output

    duplicate_summary = df_duplicate_ids.drop_duplicates(subset=["ID Auditoria"]).set_index("ID Auditoria")
    for index, row in output.iterrows():
        identifier = row.get("identificador", "")
        if identifier not in duplicate_summary.index:
            continue
        duplicate_row = duplicate_summary.loc[identifier]
        output.at[index, "DNI duplicado en entrada"] = "SI"
        output.at[index, "Riesgo cruce por DNI"] = duplicate_row.get("Riesgo cruce por DNI", "")
        output.at[index, "PDFs encontrados por DNI"] = duplicate_row.get("PDFs encontrados por nombre", "")
        output.at[index, "Certificados del DNI en entrada"] = duplicate_row.get("Numeros certificado en entrada", "")
    return output


def format_output_workbook(output_path):
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return

    workbook = load_workbook(output_path)
    input_header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    ocr_header_fill = PatternFill(fill_type="solid", fgColor="548235")
    alert_fill = PatternFill(fill_type="solid", fgColor="F4CCCC")
    ocr_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    header_font = Font(color="FFFFFF", bold=True)
    generated_columns = {
        "ID Auditoria",
        "DNI Auditoria",
        "Fecha ejecucion",
        "Lote",
        "Skip",
        "Limit",
        *USER_REVIEW_COLUMNS,
        "Cantidad apariciones",
        "Fila Excel estimada",
        "Columna ID usada",
        "Cantidad PDFs encontrados",
        "PDFs encontrados por nombre",
        "Cantidad certificados distintos",
        "Numeros certificado en entrada",
        "Filas con REVISADO informado",
        "Riesgo cruce por DNI",
        "DNI duplicado en entrada",
        "PDFs encontrados por DNI",
        "Certificados del DNI en entrada",
        "PDF auditado",
        "Fila Excel seleccionada",
        "Cantidad filas mismo DNI",
        "Estado seleccion registro",
        "Certificado usado para resolver DNI",
        "Resultado OCR",
        "Falta en archivo de entrada",
        "No encontrado en PDF",
        "Campos no validables por OCR",
    }
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = ocr_header_fill if cell.value in generated_columns else input_header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column_cells in sheet.columns:
            header = str(column_cells[0].value or "")
            width = min(max(len(header) + 2, 12), 45)
            if header in {
                "Falta en archivo de entrada",
                "No encontrado en PDF",
                "Campos no validables por OCR",
                "Comentarios",
            }:
                width = 60
            sheet.column_dimensions[column_cells[0].column_letter].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.value == "X":
                    cell.fill = alert_fill
                elif cell.value == "OCR":
                    cell.fill = ocr_fill
        if sheet.title == "Detalle CIF OCR":
            for column_letter in ["A", "D"]:
                sheet.column_dimensions[column_letter].hidden = True
            custom_widths = {
                "B": 14,
                "C": 18,
                "E": 26,
                "F": 58,
                "G": 58,
                "H": 58,
            }
            for column_letter, width in custom_widths.items():
                sheet.column_dimensions[column_letter].width = width
    workbook.save(output_path)


def ids_from_pdf_filenames(pdf_paths, pattern):
    regex = re.compile(pattern, flags=re.IGNORECASE)
    rows = []
    for pdf_path in pdf_paths:
        matches = regex.findall(pdf_path.stem)
        for match in matches:
            if isinstance(match, tuple):
                match = next((part for part in match if part), "")
            value = normalize_identifier(match)
            if value:
                rows.append(
                    {
                        "identificador": value,
                        "pagina_pdf": None,
                        "pdf": pdf_path.name,
                        "fuente": "nombre_archivo",
                    }
                )
    return rows


def preload_ocr_cache(args, pdf_dir, cache_dir):
    if cache_dir is None:
        raise ValueError("Para --solo-cache-ocr se debe indicar --cache-dir.")

    cache_dir.mkdir(parents=True, exist_ok=True)
    all_pdf_paths = sorted_pdf_paths_by_dni(pdf_dir, args.patron_id)
    pdf_paths = all_pdf_paths[args.skip_pdfs :]
    if not pdf_paths:
        raise FileNotFoundError(f"No se encontraron PDFs en {pdf_dir}")

    total = len(pdf_paths)
    cached_paths = []
    pending_paths = []
    for pdf_path in pdf_paths:
        has_cache = ocr_cache_exists(
            cache_dir,
            pdf_path,
            args.dpi,
            args.max_pages,
            args.ocr_engine,
        )
        if has_cache:
            ocr_pdf_cached(
                pdf_path=pdf_path,
                dpi=args.dpi,
                lang=args.lang,
                poppler_path=args.poppler_path,
                tesseract_cmd=args.tesseract_cmd,
                max_pages=args.max_pages,
                ocr_engine=args.ocr_engine,
                cache_dir=cache_dir,
            )
            cached_paths.append(pdf_path)
        else:
            pending_paths.append(pdf_path)

    selected_pending_paths = pending_paths
    if args.max_pdfs is not None:
        selected_pending_paths = pending_paths[: args.max_pdfs]

    processed = 0
    failed = 0
    skipped_by_limit = max(len(pending_paths) - len(selected_pending_paths), 0)
    print(f"PDFs seleccionados en carpeta: {total}", flush=True)
    print(f"Cache existente detectado: {len(cached_paths)}", flush=True)
    print(f"OCR pendientes detectados: {len(pending_paths)}", flush=True)
    if args.max_pdfs is not None:
        print(f"OCR nuevos solicitados en esta corrida: {len(selected_pending_paths)}", flush=True)
    print("", flush=True)

    for index, pdf_path in enumerate(selected_pending_paths, start=1):
        try:
            print(
                f"[{index}/{len(selected_pending_paths)}] Precargando OCR: {pdf_path.name}",
                flush=True,
            )
            ocr_pdf_cached(
                pdf_path=pdf_path,
                dpi=args.dpi,
                lang=args.lang,
                poppler_path=args.poppler_path,
                tesseract_cmd=args.tesseract_cmd,
                max_pages=args.max_pages,
                ocr_engine=args.ocr_engine,
                cache_dir=cache_dir,
            )
            processed += 1
        except Exception as exc:
            failed += 1
            print(f"ERROR OCR {pdf_path.name}: {exc}", flush=True)

    print("")
    print(f"Precarga OCR finalizada. Carpeta cache: {cache_dir.resolve()}")
    print(f"PDFs seleccionados: {total}")
    print(f"OCR nuevo generado: {processed}")
    print(f"Cache reutilizado: {len(cached_paths)}")
    print(f"Pendientes no procesados por limite: {skipped_by_limit}")
    print(f"Errores: {failed}")


def collect_ocr_rows_for_pdfs(pdf_paths, args, cache_dir, ocr_engine):
    ocr_rows = []
    raw_text_rows = []
    for pdf_path in pdf_paths:
        print(f"Leyendo/procesando OCR ({ocr_engine}): {pdf_path.name}", flush=True)
        page_texts = ocr_pdf_cached(
            pdf_path=pdf_path,
            dpi=args.dpi,
            lang=args.lang,
            poppler_path=args.poppler_path,
            tesseract_cmd=args.tesseract_cmd,
            max_pages=args.max_pages,
            ocr_engine=ocr_engine,
            cache_dir=cache_dir,
        )
        for page in page_texts:
            raw_text_rows.append(
                {
                    "pdf": pdf_path.name,
                    "pagina_pdf": page["pagina"],
                    "texto_ocr": page["texto"],
                    "motor_ocr": ocr_engine,
                }
            )
        for row in extract_matches(page_texts, args.patron_id):
            row["pdf"] = pdf_path.name
            row["fuente"] = "ocr"
            row["motor_ocr"] = ocr_engine
            ocr_rows.append(row)
    return ocr_rows, raw_text_rows


def alert_fields_by_pdf(df_field_comparison):
    fields_by_pdf = {}
    if df_field_comparison.empty:
        return fields_by_pdf
    alerts = df_field_comparison[df_field_comparison["estado"] != "OK"].copy()
    for pdf_name, group in alerts.groupby("pdf", sort=False):
        fields_by_pdf[str(pdf_name)] = sorted(set(group["campo"].astype(str)))
    return fields_by_pdf


def collect_targeted_ocr_rows_for_pdfs(pdf_paths, args, cache_dir, fields_by_pdf):
    ocr_rows = []
    raw_text_rows = []
    for pdf_path in pdf_paths:
        fields = fields_by_pdf.get(pdf_path.name, [])
        zones = retry_zones_for_fields(fields)
        if not zones:
            continue
        print(
            "Reintento OCR por zona: "
            f"{pdf_path.name} -> {', '.join(fields)} "
            f"({len(zones)} zona(s), {len(zones) * 2} llamada(s) aprox.)",
            flush=True,
        )
        page_texts = ocr_pdf_targeted_cached(
            pdf_path=pdf_path,
            dpi=args.dpi,
            fields=fields,
            lang=args.lang,
            poppler_path=args.poppler_path,
            tesseract_cmd=args.tesseract_cmd,
            max_pages=args.max_pages,
            cache_dir=cache_dir,
        )
        engine = targeted_engine_name(fields)
        for page in page_texts:
            raw_text_rows.append(
                {
                    "pdf": pdf_path.name,
                    "pagina_pdf": page["pagina"],
                    "texto_ocr": page["texto"],
                    "motor_ocr": engine,
                }
            )
        for row in extract_matches(page_texts, args.patron_id):
            row["pdf"] = pdf_path.name
            row["fuente"] = "ocr_targeted"
            row["motor_ocr"] = engine
            ocr_rows.append(row)
    return ocr_rows, raw_text_rows


def build_audit_frames(pdf_paths, df_team, df_duplicate_ids, ocr_rows, raw_text_rows, args):
    df_ocr = pd.DataFrame(ocr_rows)
    if df_ocr.empty:
        df_ocr = pd.DataFrame(columns=["identificador", "pagina_pdf", "pdf", "fuente"])

    df_pdf_ids = df_ocr.drop_duplicates(subset=["identificador"])
    df_team_ids = build_team_rows_for_pdf_selection(
        pdf_paths,
        df_team,
        raw_text_rows,
        args.patron_id,
    )
    if df_team_ids.empty:
        df_team_ids = df_team.drop_duplicates(subset=["identificador"]).copy()
    df_team_ids = add_duplicate_id_flags(df_team_ids, df_duplicate_ids)
    df_validation = build_pdf_validation(pdf_paths, df_team_ids, df_ocr, args.patron_id)
    df_field_comparison, df_cif_detail = build_field_comparison(
        pdf_paths, df_team_ids, raw_text_rows, args.patron_id
    )
    df_simple_review = build_simple_review(df_field_comparison, df_team_ids)
    return (
        df_ocr,
        df_pdf_ids,
        df_team_ids,
        df_validation,
        df_field_comparison,
        df_cif_detail,
        df_simple_review,
    )


def main():
    parser = argparse.ArgumentParser(
        description="OCR de PDFs escaneados y validacion contra un Excel/CSV del equipo."
    )
    parser.add_argument("--archivo-equipo", default=None, help="Ruta al Excel/CSV recibido.")
    parser.add_argument("--hoja", default=None, help="Hoja del Excel a leer. Si se omite usa SECLYT cuando existe.")
    parser.add_argument("--columna-id", default=None, help="Columna del archivo para cruzar.")
    parser.add_argument("--pdf-dir", required=True, help="Carpeta con PDFs escaneados.")
    parser.add_argument("--salida", default="reporte_validacion_ocr.xlsx")
    parser.add_argument("--solo-revisados", action="store_true", help="Procesa solo filas con auditoria manual.")
    parser.add_argument("--columna-revisado", default="REVISADO")
    parser.add_argument("--columna-comentarios", default="Comentarios")
    parser.add_argument(
        "--omitir-auditoria-previa",
        action="store_true",
        help="Omite filas que ya tienen un resultado humano en la columna indicada.",
    )
    parser.add_argument(
        "--columna-auditoria-previa",
        default="Auditoria - DATOS",
        help="Columna que marca que una fila ya fue auditada por una persona.",
    )
    parser.add_argument(
        "--incluir-hojas-tecnicas",
        action="store_true",
        help="Incluye hojas tecnicas de depuracion en el Excel de salida.",
    )
    parser.add_argument(
        "--consolidar",
        action="store_true",
        help="Si el Excel de salida ya existe, agrega la tanda y reemplaza DNIs repetidos.",
    )
    parser.add_argument(
        "--omitir-pdfs-procesados",
        action="store_true",
        help="Omite PDFs que ya figuran en la hoja Auditoria del Excel de salida.",
    )
    parser.add_argument(
        "--omitir-pdfs-desde-archivo",
        action="append",
        default=[],
        help="Excel adicional cuya hoja Auditoria se usa para omitir PDFs ya auditados.",
    )
    parser.add_argument(
        "--lote",
        default=None,
        help="Etiqueta del lote a guardar en el consolidado.",
    )
    parser.add_argument(
        "--dni-error-row",
        action="append",
        default=[],
        help="Correccion puntual fila_excel:DNI para celdas con error, por ejemplo 1936:92751777.",
    )
    parser.add_argument("--patron-id", default=DEFAULT_ID_PATTERN, help="Regex para extraer IDs del OCR.")
    parser.add_argument("--max-pdfs", type=int, default=None, help="Cantidad de PDFs a procesar.")
    parser.add_argument("--skip-pdfs", type=int, default=0, help="Cantidad de PDFs iniciales a saltear.")
    parser.add_argument("--max-pages", type=int, default=None, help="Maximo de paginas por PDF.")
    parser.add_argument("--dpi", type=int, default=300)
    parser.add_argument("--lang", default="spa")
    parser.add_argument("--poppler-path", default=None)
    parser.add_argument("--tesseract-cmd", default=None)
    parser.add_argument(
        "--cache-dir",
        default=None,
        help="Carpeta para guardar/reusar texto OCR por PDF.",
    )
    parser.add_argument(
        "--ocr-engine",
        choices=[
            "rapidocr",
            "rapidocr_layout",
            "rapidocr_layout_fast",
            "tesseract_pymupdf",
            "tesseract_pdf2image",
        ],
        default="rapidocr",
        help="Motor OCR. rapidocr no requiere Tesseract ni Poppler.",
    )
    parser.add_argument(
        "--reintentar-ocr-profundo",
        action="store_true",
        help="Si se usa rapidocr_layout_fast, reintenta solo las zonas de los campos con alertas.",
    )
    parser.add_argument(
        "--usar-nombre-pdf",
        action="store_true",
        help="Modo preliminar: extrae el identificador desde el nombre del PDF, sin OCR.",
    )
    parser.add_argument(
        "--solo-cache-ocr",
        action="store_true",
        help="Solo precarga el cache OCR de los PDFs seleccionados; no lee Excel ni genera reporte.",
    )
    args = parser.parse_args()

    pdf_dir = Path(args.pdf_dir)
    cache_dir = Path(args.cache_dir) if args.cache_dir else None

    if args.solo_cache_ocr:
        preload_ocr_cache(args, pdf_dir, cache_dir)
        return

    if not args.archivo_equipo:
        raise ValueError("--archivo-equipo es obligatorio salvo cuando se usa --solo-cache-ocr.")
    if not args.columna_id:
        raise ValueError("--columna-id es obligatorio salvo cuando se usa --solo-cache-ocr.")

    team_path = Path(args.archivo_equipo)

    df_team = read_team_file(team_path, sheet_name=args.hoja)
    if args.columna_id not in df_team.columns:
        available = ", ".join(map(str, df_team.columns))
        raise ValueError(f"No existe la columna '{args.columna_id}'. Columnas disponibles: {available}")

    df_team = df_team.copy()
    df_team[args.columna_id] = df_team[args.columna_id].astype("object")
    df_team["identificador"] = df_team[args.columna_id].map(normalize_identifier)
    for override in args.dni_error_row:
        if ":" not in override:
            raise ValueError("--dni-error-row debe tener formato fila_excel:DNI")
        row_number, dni = override.split(":", 1)
        dataframe_index = int(row_number) - 2
        if dataframe_index in df_team.index:
            df_team.loc[dataframe_index, args.columna_id] = dni
            df_team.loc[dataframe_index, "identificador"] = normalize_identifier(dni)
    recovered_count, identifier_mismatches = recover_missing_identifiers_from_certificate(
        df_team,
        args.columna_id,
    )
    if recovered_count:
        print(
            f"DNIs recuperados desde NumeroCertificado por celdas vacias/ilegibles: {recovered_count}",
            flush=True,
        )
    if identifier_mismatches:
        print(
            "ADVERTENCIA: hay filas donde el DNI de la columna no coincide con el DNI del NumeroCertificado. "
            "Se conserva el DNI de la columna y se recomienda revisar esas filas.",
            flush=True,
        )
        for mismatch in identifier_mismatches[:10]:
            print(
                f"  Fila {mismatch['fila_excel']}: DNI columna={mismatch['dni_columna']} "
                f"DNI certificado={mismatch['dni_certificado']}",
                flush=True,
            )
        if len(identifier_mismatches) > 10:
            print(f"  ... {len(identifier_mismatches) - 10} filas adicionales", flush=True)
    df_team = df_team[df_team["identificador"] != ""]
    df_team_duplicate_scope = df_team.copy()

    if args.omitir_auditoria_previa:
        if args.columna_auditoria_previa not in df_team.columns:
            raise ValueError(
                f"No existe la columna de auditoria previa '{args.columna_auditoria_previa}'."
            )
        previous_mask = df_team[args.columna_auditoria_previa].map(
            lambda value: normalize_manual_review(value) != ""
        )
        previous_count = int(previous_mask.sum())
        df_team = df_team[~previous_mask].copy()
        print(
            f"Registros con auditoria humana previa omitidos: {previous_count}",
            flush=True,
        )

    if args.solo_revisados:
        if args.columna_revisado not in df_team.columns:
            raise ValueError(f"No existe la columna de auditoria '{args.columna_revisado}'.")
        df_team = df_team[
            df_team[args.columna_revisado].map(normalize_manual_review) != ""
        ].copy()

    all_pdf_paths = sorted_pdf_paths_by_dni(pdf_dir, args.patron_id)
    df_duplicate_ids = build_duplicate_id_report(
        df_team_duplicate_scope,
        args.columna_id,
        all_pdf_paths,
        args.patron_id,
        args.columna_revisado,
    )
    if not df_duplicate_ids.empty:
        duplicated_count = df_duplicate_ids["ID Auditoria"].nunique()
        high_risk_count = int(
            df_duplicate_ids["Riesgo cruce por DNI"].astype(str).str.startswith("ALTO").sum()
        )
        print(
            f"ADVERTENCIA: hay {duplicated_count} DNI/ID duplicado(s) en el archivo de entrada. "
            f"Filas con riesgo alto de cruce por DNI: {high_risk_count}. "
            "Se agrega la hoja 'DNIs duplicados' al reporte.",
            flush=True,
        )
    if args.solo_revisados:
        target_ids = set(df_team["identificador"])
        all_pdf_paths = [
            pdf_path
            for pdf_path in all_pdf_paths
            if first_identifier_from_name(pdf_path, args.patron_id) in target_ids
        ]
    if args.omitir_pdfs_procesados:
        processed_pdfs = processed_pdf_names_to_omit(
            Path(args.salida),
            df_team,
            all_pdf_paths,
            args.patron_id,
            args.columna_revisado,
            args.columna_comentarios,
            solo_revisados=args.solo_revisados,
        )
        for extra_output in args.omitir_pdfs_desde_archivo:
            processed_pdfs.update(processed_pdf_names_from_output(Path(extra_output)))
        before_count = len(all_pdf_paths)
        all_pdf_paths = [
            pdf_path
            for pdf_path in all_pdf_paths
            if pdf_path.name not in processed_pdfs
        ]
        print(
            f"PDFs ya procesados omitidos: {before_count - len(all_pdf_paths)}",
            flush=True,
        )
    pdf_paths = all_pdf_paths[args.skip_pdfs :]
    if args.max_pdfs is not None:
        pdf_paths = pdf_paths[: args.max_pdfs]
    if not pdf_paths:
        if args.omitir_pdfs_procesados:
            print("")
            print("No hay PDFs pendientes para auditar con los filtros actuales.")
            print(f"Carpeta PDF: {pdf_dir}")
            print("Motivo probable: los PDFs filtrados ya figuran en el consolidado de salida.")
            print("Para reprocesar casos ya auditados, usar una corrida especifica sin --omitir-pdfs-procesados.")
            return
        raise FileNotFoundError(f"No se encontraron PDFs en {pdf_dir}")

    if args.solo_revisados:
        selected_ids = {
            first_identifier_from_name(pdf_path, args.patron_id)
            for pdf_path in pdf_paths
        }
        df_team = df_team[df_team["identificador"].isin(selected_ids)].copy()

    ocr_rows = []
    raw_text_rows = []
    if args.usar_nombre_pdf:
        print("Modo preliminar: usando identificadores desde nombres de PDF.", flush=True)
        ocr_rows = ids_from_pdf_filenames(pdf_paths, args.patron_id)
    else:
        ocr_rows, raw_text_rows = collect_ocr_rows_for_pdfs(
            pdf_paths,
            args,
            cache_dir,
            args.ocr_engine,
        )

    (
        df_ocr,
        df_pdf_ids,
        df_team_ids,
        df_validation,
        df_field_comparison,
        df_cif_detail,
        df_simple_review,
    ) = build_audit_frames(
        pdf_paths,
        df_team,
        df_duplicate_ids,
        ocr_rows,
        raw_text_rows,
        args,
    )

    if (
        args.reintentar_ocr_profundo
        and args.ocr_engine == "rapidocr_layout_fast"
        and not args.usar_nombre_pdf
        and not df_simple_review.empty
    ):
        fields_by_pdf = alert_fields_by_pdf(df_field_comparison)
        retry_pdf_paths = [
            pdf_path for pdf_path in pdf_paths if pdf_path.name in fields_by_pdf
        ]
        if retry_pdf_paths:
            print(
                f"Reintento OCR especifico por zona para {len(retry_pdf_paths)} PDF(s) con alerta.",
                flush=True,
            )
            retry_ocr_rows, retry_raw_text_rows = collect_targeted_ocr_rows_for_pdfs(
                retry_pdf_paths,
                args,
                cache_dir,
                fields_by_pdf,
            )
            ocr_rows.extend(retry_ocr_rows)
            raw_text_rows.extend(retry_raw_text_rows)
            (
                df_ocr,
                df_pdf_ids,
                df_team_ids,
                df_validation,
                df_field_comparison,
                df_cif_detail,
                df_simple_review,
            ) = build_audit_frames(
                pdf_paths,
                df_team,
                df_duplicate_ids,
                ocr_rows,
                raw_text_rows,
                args,
            )
    df_manual_comparison = compare_manual_vs_ocr(
        df_simple_review,
        df_team_ids,
        args.columna_revisado if args.solo_revisados else None,
        args.columna_comentarios if args.solo_revisados else None,
    )
    df_review_matrix = build_review_matrix(
        df_team_ids,
        df_field_comparison,
        df_simple_review,
        args.columna_revisado,
        args.columna_comentarios,
    )
    df_checkbox_view = build_checkbox_view(
        df_review_matrix,
        args.columna_revisado,
        args.columna_comentarios,
    )
    df_summary = build_summary_by_pdf(df_field_comparison)
    df_alerts = df_field_comparison[df_field_comparison["estado"] != "OK"].copy()

    present_in_pdf = set(df_pdf_ids["identificador"])
    present_in_team = set(df_team_ids["identificador"])

    df_team_ids["estado"] = df_team_ids["identificador"].map(
        lambda value: "OK_EN_PDF" if value in present_in_pdf else "FALTA_EN_PDF_O_NO_LEIDO"
    )

    only_pdf = sorted(present_in_pdf - present_in_team)
    df_only_pdf = pd.DataFrame({"identificador": only_pdf})
    if not df_only_pdf.empty:
        df_only_pdf = df_only_pdf.merge(df_pdf_ids, on="identificador", how="left")

    df_raw_text = pd.DataFrame(raw_text_rows)

    output_path = Path(args.salida)
    run_timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    lote = args.lote or f"skip_{args.skip_pdfs}_limit_{args.max_pdfs or 'all'}"
    df_checkbox_output = add_run_metadata(
        df_checkbox_view,
        run_timestamp,
        lote,
        args.skip_pdfs,
        args.max_pdfs,
    )
    df_checkbox_output = ensure_user_review_columns(df_checkbox_output)
    df_alerts_output = add_run_metadata(
        df_alerts,
        run_timestamp,
        lote,
        args.skip_pdfs,
        args.max_pdfs,
    )
    df_cif_detail_output = add_run_metadata(
        df_cif_detail,
        run_timestamp,
        lote,
        args.skip_pdfs,
        args.max_pdfs,
    )
    df_cif_detail_output = order_cif_detail_columns(df_cif_detail_output)
    if args.consolidar:
        df_checkbox_output = merge_consolidated_sheet(
          output_path,
          "Auditoria",
          df_checkbox_output,
          ["ID Auditoria"],
          USER_REVIEW_COLUMNS,
      )
        df_alerts_output = merge_consolidated_sheet(
            output_path,
            "Detalle alertas",
            df_alerts_output,
            ["ID Auditoria", "campo"],
        )
        df_cif_detail_output = merge_consolidated_sheet(
            output_path,
            "Detalle CIF OCR",
            df_cif_detail_output,
            ["ID Auditoria", "campo"],
        )
        df_cif_detail_output = order_cif_detail_columns(df_cif_detail_output)

    with pd.ExcelWriter(output_path) as writer:
        df_checkbox_output.to_excel(writer, sheet_name="Auditoria", index=False)
        df_alerts_output.to_excel(writer, sheet_name="Detalle alertas", index=False)
        df_cif_detail_output.to_excel(writer, sheet_name="Detalle CIF OCR", index=False)
        if not df_duplicate_ids.empty:
            df_duplicate_ids.to_excel(writer, sheet_name="DNIs duplicados", index=False)
        if args.incluir_hojas_tecnicas:
            df_review_matrix.to_excel(writer, sheet_name="misma_estructura", index=False)
            if not df_manual_comparison.empty:
                df_manual_comparison.to_excel(writer, sheet_name="auditoria_manual_vs_ocr", index=False)
            df_simple_review.to_excel(writer, sheet_name="revision_simple", index=False)
            df_summary.to_excel(writer, sheet_name="resumen_por_pdf", index=False)
            df_validation.to_excel(writer, sheet_name="validacion_por_pdf", index=False)
            df_field_comparison.to_excel(writer, sheet_name="comparacion_campos", index=False)
            df_team_ids.to_excel(writer, sheet_name="archivo_equipo_validado", index=False)
            df_ocr.to_excel(writer, sheet_name="ids_detectados_pdf", index=False)
            df_only_pdf.to_excel(writer, sheet_name="en_pdf_no_en_archivo", index=False)
            df_raw_text.to_excel(writer, sheet_name="texto_ocr_paginas", index=False)

    format_output_workbook(output_path)

    print(f"Reporte generado: {output_path.resolve()}")
    print(f"Registros del archivo: {len(df_team_ids)}")
    print(f"IDs detectados en PDFs: {len(df_pdf_ids)}")
    print(f"En archivo y encontrados en PDF: {len(present_in_team & present_in_pdf)}")
    print(f"En archivo pero no encontrados: {len(present_in_team - present_in_pdf)}")
    print(f"En PDF pero no en archivo: {len(present_in_pdf - present_in_team)}")
    if not df_simple_review.empty:
        print("Resumen revision simple:")
        for result, count in df_simple_review["resultado"].value_counts().items():
            print(f"  {result}: {count}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        sys.exit(1)
